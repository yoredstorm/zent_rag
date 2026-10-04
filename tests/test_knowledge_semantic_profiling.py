# =============================================================================
# Semantic Column Profiling — tests (§17)
# =============================================================================
# Fechas, IDs, enums de baja cardinalidad, claves candidatas, códigos y
# unidades. Todo determinista y sin dominios hardcodeados en el core.
# =============================================================================
from __future__ import annotations

import io
from uuid import UUID

from openpyxl import Workbook

from src.core.domain.tabular import TabularSemanticType, TabularValueType
from src.knowledge.tabular.builder import build_tabular_workbook
from src.knowledge.tabular.schema import (
    infer_format,
    infer_units,
    infer_value_pattern,
    is_candidate_key,
    likely_foreign_key,
)
from src.knowledge.tabular.xlsx_reader import read_xlsx_grid

_ORG = UUID("00000000-0000-0000-0000-0000000000cc")


def test_infer_format_detects_yyyymmdd() -> None:
    fmt, confidence = infer_format(
        ("20240115", "20240201", "20241231"),
        value_type=TabularValueType.INTEGER,
    )
    assert fmt == "YYYYMMDD"
    assert confidence > 0.7


def test_infer_format_detects_iso() -> None:
    fmt, _confidence = infer_format(
        ("2024-01-15", "2024-02-01"),
        value_type=TabularValueType.DATE,
    )
    assert fmt == "YYYY-MM-DD"


def test_infer_format_marks_ambiguous_slash_dates() -> None:
    fmt, confidence = infer_format(
        ("01/02/2024", "03/04/2024"),
        value_type=TabularValueType.STRING,
    )
    assert fmt == "DD/MM/YYYY"
    # Ambigua: no se decide con confianza alta.
    assert confidence <= 0.65


def test_infer_format_rejects_invalid_compact_dates() -> None:
    fmt, _confidence = infer_format(
        ("99999999", "12345678"),
        value_type=TabularValueType.INTEGER,
    )
    assert fmt is None


def test_infer_value_pattern() -> None:
    assert infer_value_pattern(("CAT31", "CAT32", "CAT33")) == "AAA99"
    assert infer_value_pattern(("CXRCD", "ABCDF")) == "AAAAA"
    assert infer_value_pattern(("uno", "dos", "tres")) is None


def test_infer_units() -> None:
    assert infer_units(("150.50 USD", "200 USD", "10 USD"))[0] == "USD"
    assert infer_units(("5%", "10%", "12.5%"))[0] == "%"
    assert infer_units(("hola", "mundo"))[0] is None


def test_candidate_key_requires_not_nullable_and_unique() -> None:
    assert is_candidate_key(non_empty=3, unique_ratio=1.0, nullable=False) is True
    assert is_candidate_key(non_empty=3, unique_ratio=1.0, nullable=True) is False
    assert is_candidate_key(non_empty=3, unique_ratio=0.5, nullable=False) is False


def test_likely_foreign_key_signal() -> None:
    fk, confidence = likely_foreign_key(
        "CUSTOMER_ID",
        TabularSemanticType.IDENTIFIER,
        candidate_key=False,
        unique_ratio=0.4,
    )
    assert fk is True
    assert confidence > 0.5
    # Una clave candidata no es FK.
    assert likely_foreign_key(
        "CUSTOMER_ID",
        TabularSemanticType.IDENTIFIER,
        candidate_key=True,
        unique_ratio=1.0,
    ) == (False, 0.0)
    assert likely_foreign_key(
        "DESCRIPTION",
        TabularSemanticType.DESCRIPTION,
        candidate_key=False,
        unique_ratio=0.1,
    ) == (False, 0.0)


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    headers = ["FPROC", "CCUST", "CARRIER_CODE", "AMOUNT", "STATUS"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=column, value=header)
    rows = [
        ["20240115", "C001", "AA", "150.50 USD", "A"],
        ["20240201", "C002", "AM", "200.00 USD", "A"],
        ["20241231", "C003", "AA", "99.99 USD", "I"],
    ]
    for row_index, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_generic_value_fallbacks_without_domain_vocabulary() -> None:
    """Email, %, código corto y entero único: heurísticas por valor, no dominio."""
    from src.knowledge.tabular.schema import infer_semantic_type

    assert infer_semantic_type(
        "CONTACT",
        value_type=TabularValueType.STRING,
        samples=("ana@example.com", "luis@example.com"),
        unique_ratio=1.0,
    )[0] is TabularSemanticType.EMAIL
    assert infer_semantic_type(
        "DISCOUNT",
        value_type=TabularValueType.STRING,
        samples=("5%", "10%", "12.5%"),
        unique_ratio=0.5,
    )[0] is TabularSemanticType.PERCENTAGE
    assert infer_semantic_type(
        "CARR",
        value_type=TabularValueType.STRING,
        samples=("AA", "AM", "AA"),
        unique_ratio=0.66,
    )[0] is TabularSemanticType.CODE
    assert infer_semantic_type(
        "SEQ_NO",
        value_type=TabularValueType.INTEGER,
        max_length=6,
        samples=("1001", "1002", "1003"),
        unique_ratio=1.0,
    )[0] is TabularSemanticType.IDENTIFIER
    assert infer_semantic_type(
        "TOTAL",
        value_type=TabularValueType.STRING,
        samples=("$10.50", "€20.00"),
        unique_ratio=0.5,
    )[0] is TabularSemanticType.AMOUNT
    # Header fuerte gana sobre el fallback por valor (no se pisa semántica).
    assert infer_semantic_type(
        "Start Position",
        value_type=TabularValueType.INTEGER,
        samples=("1", "2"),
        unique_ratio=1.0,
    )[0] is TabularSemanticType.POSITION


def test_tabular_adversarial_table_inferences_are_reasonable() -> None:
    """Tabla completa: fechas, códigos, IDs, porcentajes, texto largo."""
    from src.knowledge.tabular.schema import infer_format, infer_semantic_type

    # DATE_COLUMN: 20260101/20260102 -> YYYYMMDD, no ID genérico.
    fmt, confidence = infer_format(
        ("20260101", "20260102"), value_type=TabularValueType.INTEGER
    )
    assert fmt == "YYYYMMDD"
    assert confidence > 0.7
    # CODE genérico: no se asume aerolínea ni dominio.
    semantic, _confidence, _aliases = infer_semantic_type(
        "CODE",
        value_type=TabularValueType.STRING,
        samples=("AM", "AA", "LA"),
        unique_ratio=1.0,
    )
    assert semantic is TabularSemanticType.CODE
    # Adversarial: códigos arbitrarios siguen siendo CODE genérico, no dominio.
    semantic, _confidence, aliases = infer_semantic_type(
        "CODE",
        value_type=TabularValueType.STRING,
        samples=("AB", "CD", "EF"),
        unique_ratio=1.0,
    )
    assert semantic is TabularSemanticType.CODE
    assert "airline" not in aliases
    # ID 1,2,3 -> identificador.
    assert infer_semantic_type(
        "ID",
        value_type=TabularValueType.INTEGER,
        samples=("1", "2", "3"),
        unique_ratio=1.0,
    )[0] is TabularSemanticType.IDENTIFIER
    # PERCENT 0.12/0.20 -> porcentaje por header.
    assert infer_semantic_type(
        "PERCENT",
        value_type=TabularValueType.FLOAT,
        samples=("0.12", "0.20"),
        unique_ratio=0.5,
    )[0] is TabularSemanticType.PERCENTAGE
    # TEXT largo -> free text.
    assert infer_semantic_type(
        "TEXT",
        value_type=TabularValueType.STRING,
        max_length=240,
        samples=("descripcion larga " * 10,),
        unique_ratio=1.0,
    )[0] is TabularSemanticType.FREE_TEXT


def test_builder_enriches_columns_with_semantic_profile() -> None:
    grid = read_xlsx_grid(_workbook_bytes(), "profiling.xlsx")
    workbook = build_tabular_workbook(
        grid, organization_id=_ORG, external_id="profiling.xlsx"
    )
    table = workbook.tables()[0]
    by_name = {column.original_name: column for column in table.columns}

    fproc = by_name["FPROC"]
    assert fproc.metadata["format"] == "YYYYMMDD"
    assert fproc.metadata["value_pattern"] == "99999999"
    assert fproc.metadata["candidate_key"] is True

    ccust = by_name["CCUST"]
    assert ccust.metadata["value_pattern"] == "A999"
    assert ccust.metadata["candidate_key"] is True

    carrier = by_name["CARRIER_CODE"]
    assert carrier.metadata["value_pattern"] == "AA"
    assert carrier.metadata["candidate_key"] is False
    assert carrier.metadata["likely_foreign_key"] is True

    amount = by_name["AMOUNT"]
    assert amount.metadata["units"] == "USD"

    # Nada inventado: columnas sin señal quedan None/False.
    status = by_name["STATUS"]
    assert status.metadata["candidate_key"] is False
    assert status.metadata["format"] is None
    assert len(status.metadata["examples_bounded"]) <= 8
