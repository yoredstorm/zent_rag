# =============================================================================
# Knowledge Compiler — pruebas reales con PDF, Excel y CSV
# =============================================================================
# Nada de mocks del parser: se construyen archivos reales en memoria y se
# pasan por los parsers productivos (PdfParser, XlsxParser, CsvParser).
#
# Lo que se verifica:
#   1. units -> entities -> facts -> relationships -> rules sin pérdida
#   2. identidad canónica compartida entre archivos (misma entidad, una sola)
#   3. provenance completa: documento/página/sección/bloque y tabla/fila/celda
#   4. temporalidad, conflictos clasificados y corrobación (no duplicación)
# =============================================================================
from __future__ import annotations

import zlib
from datetime import date
from uuid import UUID, uuid4

import pytest

from src.core.domain.canonical import CanonicalKind, canonical_uuid
from src.core.domain.knowledge_events import (
    KnowledgeEventType,
    KnowledgeSystemEvent,
)
from src.core.domain.knowledge_v2 import StructuredDocument
from src.knowledge.compiler import (
    CompilationResult,
    ConflictCandidate,
    ConflictType,
    EntityResolver,
    KnowledgeCompiler,
    RuleCandidate,
    SemanticUnitKind,
    classify_conflict,
    detect_conflicts,
    discover_entities,
    extract_rules,
    extract_semantic_units,
    extract_tabular_units,
    infer_temporal_scope,
    values_equivalent,
)
from src.knowledge.compiler.model import FactCandidate, FactKind, TemporalScope
from src.knowledge.structure.csv_parser import CsvParser
from src.knowledge.structure.pdf_parser import PdfParser
from src.knowledge.structure.xlsx_parser import XlsxParser

ORG = UUID("11111111-2222-3333-4444-555555555555")
SOURCE_A = UUID("aaaaaaaa-0000-0000-0000-000000000001")
SOURCE_B = UUID("bbbbbbbb-0000-0000-0000-000000000002")


# ---------------------------------------------------------------------------
# Archivos reales
# ---------------------------------------------------------------------------


def _pdf(pages: list[list[str]]) -> bytes:
    """PDF mínimo real de varias páginas (mismo generador que los tests de DU)."""
    streams = [zlib.compress("\n".join(ops).encode("latin-1")) for ops in pages]
    count = len(pages)
    font_id = 3 + 2 * count
    objects: list[bytes] = []
    kids = " ".join(f"{3 + 2 * index} 0 R" for index in range(count))
    objects.append(b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj")
    objects.append(
        f"2 0 obj << /Type /Pages /Kids [{kids}] /Count {count} >> endobj".encode()
    )
    for index, stream in enumerate(streams):
        page_id = 3 + 2 * index
        content_id = page_id + 1
        objects.append(
            (
                f"{page_id} 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >> endobj"
            ).encode()
        )
        objects.append(
            f"{content_id} 0 obj << /Length {len(stream)} /Filter /FlateDecode >> stream\n".encode()
            + stream
            + b"\nendstream endobj"
        )
    objects.append(
        f"{font_id} 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj".encode()
    )
    body = b"\n".join(objects)
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for obj in objects:
        offsets.append(len(out))
        out += obj + b"\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF"
    ).encode()
    return bytes(out)


def _line(text: str, x: float, y: float) -> str:
    return f"BT /F1 11 Tf {x} {y} Td ({text}) Tj ET"


def record4_pdf() -> bytes:
    """Manual real: definiciones, campo técnico, regla y procedimiento."""
    page_one = [
        _line("ATPCO Record 4 (R4)", 72, 700),
        _line("Record 4: renumbers the Record 2 sequence number on reissue.", 72, 670),
        _line("Byte 105 | 3 | Currency Code | ISO 4217 currency of the fare component", 72, 640),
        _line("Effective from 01 March 2025, revision 3.", 72, 610),
        _line("The carrier must not change the fare basis code on reissue.", 72, 580),
    ]
    page_two = [
        _line("Cat 31 Voluntary Changes", 72, 700),
        _line("Category 31 (Cat 31): voluntary changes to a ticket.", 72, 670),
        _line("The carrier must reissue the ticket within 24 hours.", 72, 640),
    ]
    return _pdf([page_one, page_two])


def record2_pdf() -> bytes:
    """Segundo manual, otra fuente, mismo dominio."""
    page_one = [
        _line("ATPCO Record 2", 72, 700),
        _line("Record 2: sequence number assigned to each fare component.", 72, 670),
        _line("Effective from 01 June 2026, revision 4.", 72, 640),
        _line("Record 4: renumbers the Record 2 sequence number on reissue.", 72, 610),
    ]
    return _pdf([page_one])


def fare_table_xlsx() -> bytes:
    """Excel real: columnas tipadas, con la entidad Record 4 en la cabecera."""
    import io

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Record4"
    sheet["A1"] = "ATPCO RECORD 4 MAPPING"
    headers = ["Record 4", "Category 31", "Byte", "Meaning"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=column, value=header)
    rows = [
        ["R4-1", "Cat 31", 105, "Currency code of the fare component"],
        ["R4-2", "Cat 31", 106, "Fare basis code"],
    ]
    for row_index, row in enumerate(rows, start=4):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def carriers_csv() -> bytes:
    return (
        "Carrier Code,Category,Meaning\n"
        "AA,Cat 31,Voluntary changes\n"
        "UA,Category 31,Voluntary changes\n"
    ).encode("utf-8")


def parse_pdf(data: bytes, *, external_id: str, source_id: UUID) -> StructuredDocument:
    from src.knowledge.structure.pdf_parser import PdfParseOptions
    from src.knowledge.understanding.engine import apply_understanding

    document = PdfParser().parse(
        data,
        organization_id=ORG,
        external_id=external_id,
        source_id=source_id,
        source_name=external_id,
        options=PdfParseOptions(column_detection=False),
    )
    return apply_understanding(document, filename=external_id)


def parse_xlsx(data: bytes, *, external_id: str, source_id: UUID) -> StructuredDocument:
    from src.knowledge.understanding.engine import apply_understanding

    document = XlsxParser().parse(
        data,
        organization_id=ORG,
        external_id=external_id,
        source_id=source_id,
        source_name=external_id,
    )
    return apply_understanding(document, filename=external_id)


def parse_csv(data: bytes, *, external_id: str, source_id: UUID) -> StructuredDocument:
    from src.knowledge.understanding.engine import apply_understanding

    document = CsvParser().parse(
        data,
        organization_id=ORG,
        external_id=external_id,
        source_id=SOURCE_B,
        source_name=external_id,
    )
    return apply_understanding(document, filename=external_id)


# ---------------------------------------------------------------------------
# 1. Unidades semánticas con provenance
# ---------------------------------------------------------------------------


def test_pdf_real_produce_unidades_con_locator_completo() -> None:
    document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)
    units = extract_semantic_units(document)

    definitions = [u for u in units if u.kind == SemanticUnitKind.DEFINITION.value]
    assert any(u.label == "Record 4" for u in definitions)
    record4 = next(u for u in definitions if u.label == "Record 4")
    assert record4.evidence.locator.document_id == document.id
    assert record4.evidence.locator.page == 1
    assert record4.evidence.locator.block_id is not None
    assert record4.evidence.locator.locator_uri().startswith(f"document/{document.id}")

    fields = [u for u in units if u.kind == SemanticUnitKind.FIELD.value]
    byte_unit = next(u for u in fields if u.label.startswith("Byte 105"))
    assert byte_unit.attributes["length"] == 3
    assert byte_unit.evidence.locator.page == 1

    # La segunda página también aporta unidades, con su propia página.
    assert any(
        u.evidence.locator.page == 2 for u in units if u.evidence.locator.page
    )


def test_reglas_solo_desde_lenguaje_normativo() -> None:
    document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)
    rules = extract_rules(document)

    statements = [rule.statement for rule in rules]
    assert any("must not change the fare basis code" in s for s in statements)
    assert any("must reissue the ticket" in s for s in statements)
    # Una descripción sin modalidad normativa no es una regla.
    assert not any("renumbers the Record 2 sequence number" == s for s in statements)
    assert {rule.modality for rule in rules} <= {"must", "must_not", "constraint", "should"}
    for rule in rules:
        assert rule.evidence, "regla sin evidencia"
        assert rule.evidence[0].locator.document_id == document.id


def test_excel_real_expone_estructura_no_texto() -> None:
    document = parse_xlsx(fare_table_xlsx(), external_id="record4.xlsx", source_id=SOURCE_A)
    units = extract_tabular_units(document.tabular, document=document)

    tables = [u for u in units if u.kind == SemanticUnitKind.TABLE.value]
    assert tables, "no se detectó la tabla del workbook"
    table = tables[0]
    assert table.evidence.locator.table_reference
    assert table.attributes["row_count"] == 2
    assert table.attributes["column_count"] == 4

    columns = [u for u in units if u.kind == SemanticUnitKind.COLUMN.value]
    names = {u.label for u in columns}
    assert {"Record 4", "Category 31", "Byte", "Meaning"}.issubset(names)
    byte_column = next(u for u in columns if u.label == "Byte")
    assert byte_column.evidence.locator.table_reference == table.label
    assert byte_column.evidence.locator.cell_reference == "C1"
    assert byte_column.attributes["inferred_type"] in {"int", "integer"}


def test_csv_real_expone_columnas_y_tipos() -> None:
    document = parse_csv(carriers_csv(), external_id="carriers.csv", source_id=SOURCE_B)
    units = extract_tabular_units(document.tabular, document=document)

    columns = {u.label for u in units if u.kind == SemanticUnitKind.COLUMN.value}
    assert {"Carrier Code", "Category", "Meaning"}.issubset(columns)


# ---------------------------------------------------------------------------
# 2. Identidad canónica entre archivos
# ---------------------------------------------------------------------------


def test_misma_entidad_en_pdf_y_excel_comparte_identidad_canonica() -> None:
    pdf_document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)
    xlsx_document = parse_xlsx(
        fare_table_xlsx(), external_id="record4.xlsx", source_id=SOURCE_A
    )

    pdf_entities = discover_entities(extract_semantic_units(pdf_document), document=pdf_document)
    xlsx_entities = discover_entities(
        extract_semantic_units(xlsx_document)
        + extract_tabular_units(xlsx_document.tabular, document=xlsx_document),
        document=xlsx_document,
    )

    pdf_record4 = next(e for e in pdf_entities if e.name == "Record 4")
    xlsx_record4 = next(e for e in xlsx_entities if e.name == "Record 4")

    # La identidad es determinista por (organización, tipo lógico, clave natural).
    assert pdf_record4.natural_key == xlsx_record4.natural_key
    kind = "entity"
    assert canonical_uuid(ORG, CanonicalKind(kind), pdf_record4.natural_key) == canonical_uuid(
        ORG, CanonicalKind(kind), xlsx_record4.natural_key
    )
    # Y la evidencia es distinta: una viene del PDF, otra del Excel.
    pdf_locator = pdf_record4.evidence[0].locator
    xlsx_locator = xlsx_record4.evidence[0].locator
    assert pdf_locator.document_id != xlsx_locator.document_id
    assert xlsx_locator.table_reference


def test_alias_declarado_fusiona_con_razon_y_confianza() -> None:
    document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)
    entities, merges = EntityResolver().consolidate(
        discover_entities(extract_semantic_units(document), document=document)
    )

    record4 = next(e for e in entities if e.name == "Record 4")
    assert any(alias.alias == "R4" for alias in record4.aliases)
    r4 = next(alias for alias in record4.aliases if alias.alias == "R4")
    assert r4.confidence >= 0.85
    assert r4.reason, "alias sin razón de merge"
    assert r4.alias_type in {"acronym", "abbreviation", "synonym"}


def test_entidades_sin_evidencia_no_se_fusionan() -> None:
    """Record 4 y Record 2 son entidades distintas: no se unen por parecido."""
    document = parse_pdf(record2_pdf(), external_id="record2.pdf", source_id=SOURCE_B)
    entities, merges = EntityResolver().consolidate(
        discover_entities(extract_semantic_units(document), document=document)
    )
    names = {e.name for e in entities}
    assert "Record 2" in names
    assert "Record 4" in names
    assert all(
        merge.canonical_name != merge.merged_alias for merge in merges
    ), "merge sin cambio de nombre"
    for merge in merges:
        assert merge.reason and merge.confidence > 0


# ---------------------------------------------------------------------------
# 3. Hechos, relaciones, temporalidad y conflictos
# ---------------------------------------------------------------------------


def test_hechos_conservan_sujeto_predicado_y_evidencia() -> None:
    document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)
    result = KnowledgeCompiler.build(document)

    definition_fact = next(
        fact
        for fact in result.facts
        if fact.subject == "Record 4" and fact.predicate == "defined_as"
    )
    assert "renumbers the Record 2 sequence number" in (definition_fact.object_value or "")
    assert definition_fact.fact_kind == FactKind.DEFINITION.value
    assert definition_fact.evidence
    assert definition_fact.evidence[0].locator.page == 1

    field_facts = {
        (fact.predicate, fact.object_value)
        for fact in result.facts
        if fact.subject.startswith("Byte 105")
    }
    assert ("has_length", "3") in field_facts


def test_relaciones_estructurales_de_columnas() -> None:
    document = parse_xlsx(fare_table_xlsx(), external_id="record4.xlsx", source_id=SOURCE_A)
    result = KnowledgeCompiler.build(document)

    pairs = {(r.subject, r.predicate, r.object_name) for r in result.relationships}
    assert any(
        subject == "Byte" and predicate == "part_of"
        for subject, predicate, _object in pairs
    )
    for relationship in result.relationships:
        assert relationship.evidence, "relación sin evidencia"


def test_vigencia_declarada_queda_acotada() -> None:
    document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)
    scope = infer_temporal_scope(document)
    assert scope.effective_from == date(2025, 3, 1)
    assert scope.version_label == "3"
    assert scope.is_bounded


def test_mismo_sujeto_con_vigencias_disjuntas_es_cambio_temporal() -> None:
    fact_a = FactCandidate(
        subject="Record 4",
        predicate="has_length",
        object_value="3",
        temporal=TemporalScope(effective_from=date(2025, 1, 1), effective_to=date(2025, 12, 31)),
        evidence=[],
    )
    fact_b = FactCandidate(
        subject="Record 4",
        predicate="has_length",
        object_value="4",
        temporal=TemporalScope(effective_from=date(2026, 1, 1)),
        evidence=[],
    )
    conflict = classify_conflict(fact_a, fact_b)
    assert conflict.conflict_type == ConflictType.TEMPORAL_CHANGE.value
    assert conflict.confidence > 0.5
    assert "evolución temporal" in conflict.reason


def test_mismo_valor_escrito_distinto_es_posible_duplicado() -> None:
    assert values_equivalent("Record 4", "record 4")
    assert values_equivalent("ATPCO Record 4", "Record 4")
    assert not values_equivalent("Record 4", "Record 2")
    fact_a = FactCandidate(subject="Cat 31", predicate="meaning", object_value="Voluntary Changes")
    fact_b = FactCandidate(subject="Cat 31", predicate="meaning", object_value="voluntary  changes")
    conflict = classify_conflict(fact_a, fact_b)
    assert conflict.conflict_type == ConflictType.POSSIBLE_DUPLICATE.value
    assert conflict.values_equivalent is True


def test_conflicto_entre_fuentes_independientes() -> None:
    fact_a = FactCandidate(
        subject="Byte 105",
        predicate="has_length",
        object_value="3",
        attributes={"source_label": "ATPCO Record 4"},
    )
    fact_b = FactCandidate(
        subject="Byte 105",
        predicate="has_length",
        object_value="5",
        attributes={"source_label": "Carrier Manual"},
    )
    conflict = classify_conflict(fact_a, fact_b)
    assert conflict.conflict_type == ConflictType.SOURCE_CONFLICT.value
    assert conflict.source_a == "ATPCO Record 4"
    assert conflict.source_b == "Carrier Manual"


def test_detect_conflicts_no_reporta_hechos_identicos() -> None:
    base = dict(subject="Record 4", predicate="has_length")
    same_a = FactCandidate(**base, object_value="3")
    same_b = FactCandidate(**base, object_value="3")
    assert detect_conflicts([same_a, same_b]) == []
    differing = FactCandidate(**base, object_value="4")
    conflicts = detect_conflicts([same_a, differing])
    assert len(conflicts) == 1
    assert conflicts[0].subject == "Record 4"


def test_compilacion_completa_de_pdf_real() -> None:
    document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)
    result = KnowledgeCompiler.build(document)

    counts = result.to_dict()["counts"]
    assert counts["units"] > 0
    assert counts["entities"] > 0
    assert counts["facts"] > 0
    assert counts["rules"] >= 2
    assert counts["evidence"] > 0
    assert result.document_id == document.id
    assert result.source_id == SOURCE_A

    # Toda evidencia apunta al documento de la fuente que la produjo.
    for unit in result.units:
        if unit.evidence.locator.document_id is not None:
            assert unit.evidence.locator.document_id == document.id


# ---------------------------------------------------------------------------
# 4. Persistencia: corrobación sin duplicación y provenance enlazada
# ---------------------------------------------------------------------------


class FakeCompilerStore:
    """Store en memoria con la misma semántica de corrobación del real."""

    def __init__(self) -> None:
        self.objects: dict[str, dict] = {}
        self.aliases: dict[str, tuple[str, str]] = {}
        self.facts: dict[tuple[str, str, str], dict] = {}
        self.edges: dict[tuple[str, str, str], dict] = {}
        self.rules: dict[str, dict] = {}
        self.evidence: list[dict] = []
        self.conflicts: list[dict] = []
        self.compilations: list[dict] = []

    async def existing_aliases(self, organization_id):
        return {alias: canonical for alias, (_id, canonical) in self.aliases.items()}

    async def upsert_entity(self, organization_id, entity, *, source_id, workspace_id):
        object_id = canonical_uuid(ORG, CanonicalKind("entity"), entity.natural_key)
        record = self.objects.setdefault(
            str(object_id),
            {"name": entity.name, "evidence": 0, "sources": set()},
        )
        record["evidence"] += len(entity.evidence)
        if source_id:
            record["sources"].add(str(source_id))
        self.aliases[entity.normalized] = (str(object_id), entity.name)
        return object_id

    async def upsert_alias(
        self, organization_id, *, entity_id, alias, source_id, document_id, evidence_ids
    ):
        self.aliases[alias.normalized] = (str(entity_id), self.aliases.get(alias.normalized, (None, alias.alias))[1])

    async def upsert_fact(
        self, organization_id, fact, *, subject_id, source_id, workspace_id, document_id
    ):
        key = (fact.subject, fact.predicate, fact.object_value or "")
        record = self.facts.get(key)
        created = record is None
        if record is None:
            record = {
                "subject_id": subject_id,
                "sources": set(),
                "evidence": 0,
                "corroborated": False,
                "count": 1,
            }
            self.facts[key] = record
        else:
            record["count"] += 1
        if source_id:
            if record["sources"] and str(source_id) not in record["sources"]:
                record["corroborated"] = True
            record["sources"].add(str(source_id))
        record["evidence"] += len(fact.evidence)
        status = "created" if created else "reinforced"
        return uuid4(), len(fact.evidence), status

    async def upsert_relationship(
        self, organization_id, relationship, *, subject_id, object_id, source_id, workspace_id
    ):
        if subject_id is None or object_id is None:
            return None, "skipped"
        key = (str(subject_id), relationship.predicate, str(object_id))
        created = key not in self.edges
        self.edges[key] = {
            "confidence": relationship.confidence,
            "evidence": len(relationship.evidence),
        }
        return uuid4(), ("created" if created else "reinforced")

    async def upsert_rule(self, organization_id, rule, *, source_id, workspace_id, document_id):
        object_id = canonical_uuid(ORG, CanonicalKind.BUSINESS_RULE, f"rule:{rule.rule_key}")
        created = str(object_id) not in self.rules
        self.rules[str(object_id)] = {
            "statement": rule.statement,
            "subject": rule.subject,
            "rule_key": rule.rule_key,
            "evidence": len(rule.evidence),
            "sources": {str(source_id)} if source_id else set(),
        }
        return object_id, ("created" if created else "reinforced")

    async def existing_rule_keys(self, organization_id):
        return {row["subject"]: row["rule_key"] for row in self.rules.values()}

    async def add_evidence(
        self, organization_id, evidence, *, canonical_id=None, assertion_id=None, workspace_id=None, authority=None
    ):
        self.evidence.append(
            {
                "canonical_id": str(canonical_id) if canonical_id else None,
                "assertion_id": str(assertion_id) if assertion_id else None,
                "locator": evidence.locator.locator_uri(),
                "document_id": str(evidence.locator.document_id)
                if evidence.locator.document_id
                else None,
                "page": evidence.locator.page,
                "table": evidence.locator.table_reference,
                "cell": evidence.locator.cell_reference,
                "row": evidence.locator.row_reference,
                "excerpt": evidence.excerpt,
                "strength": evidence.strength,
                "evidence_type": evidence.evidence_type,
            }
        )
        return uuid4()

    async def upsert_conflict(self, organization_id, conflict, *, object_id, workspace_id):
        duplicated = any(
            existing["subject"] == conflict.subject
            and existing["type"] == conflict.conflict_type
            and existing["reason"] == conflict.reason
            for existing in self.conflicts
        )
        self.conflicts.append(
            {
                "type": conflict.conflict_type,
                "reason": conflict.reason,
                "subject": conflict.subject,
            }
        )
        return "duplicate" if duplicated else "created"

    async def record_compilation(
        self, organization_id, *, result, workspace_id, duration_ms, status, error=None
    ):
        self.compilations.append({"status": status, "error": error, **result.to_dict()["counts"]})

    async def refresh_counters(self, organization_id) -> None:
        return None


class FakeSystemEmitter:
    """Emisor de eventos de sistema en memoria: registra lo emitido."""

    def __init__(self) -> None:
        self.events: list[KnowledgeSystemEvent] = []

    async def emit(self, event: KnowledgeSystemEvent) -> None:
        self.events.append(event)

    def of_type(self, event_type: KnowledgeEventType) -> list[KnowledgeSystemEvent]:
        return [event for event in self.events if event.type == event_type]


class PreparedCompiler(KnowledgeCompiler):
    """Compiler con resultado determinista: aísla el mapeo del pipeline."""

    def __init__(self, result: CompilationResult, *, store=None) -> None:
        super().__init__(store=store)
        self._result = result

    def build(self, document: StructuredDocument) -> CompilationResult:  # type: ignore[override]
        return self._result


def _prepared_result(**overrides) -> CompilationResult:
    base = {
        "organization_id": ORG,
        "source_id": SOURCE_A,
        "document_id": uuid4(),
        "document_title": "manual.pdf",
    }
    base.update(overrides)
    return CompilationResult(**base)


@pytest.mark.asyncio
async def test_dos_fuentes_corroboran_el_mismo_hecho_sin_duplicarlo() -> None:
    store = FakeCompilerStore()
    compiler = KnowledgeCompiler(store=store)

    first = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)
    second = parse_pdf(record2_pdf(), external_id="record2.pdf", source_id=SOURCE_B)

    await compiler.compile_document(first)
    assert store.facts, "no se persistió ningún hecho"

    await compiler.compile_document(second)

    shared_key = next(
        key
        for key in store.facts
        if key[0] == "Record 4" and key[1] == "defined_as"
    )
    record = store.facts[shared_key]
    assert record["count"] >= 2, "el mismo hecho se perdió al re-compilar"
    assert str(SOURCE_A) in record["sources"]
    assert str(SOURCE_B) in record["sources"]
    assert record["corroborated"] is True
    assert len(store.facts) == len(
        {key for key in store.facts}
    ), "los hechos se duplicaron en lugar de reforzarse"


@pytest.mark.asyncio
async def test_provenance_sobrevive_hasta_la_celda() -> None:
    store = FakeCompilerStore()
    compiler = KnowledgeCompiler(store=store)

    document = parse_xlsx(fare_table_xlsx(), external_id="record4.xlsx", source_id=SOURCE_A)
    await compiler.compile_document(document)

    assert store.evidence, "el compilador no escribió evidencia"
    cell_evidence = [item for item in store.evidence if item["cell"]]
    assert cell_evidence, "no hay evidencia con referencia de celda"
    for item in store.evidence:
        assert item["locator"]
        assert item["document_id"] == str(document.id)
    table_rows = [item for item in store.evidence if item["table"]]
    assert table_rows, "no hay evidencia con referencia de tabla"


@pytest.mark.asyncio
async def test_compilacion_registra_traza_y_es_idempotente() -> None:
    store = FakeCompilerStore()
    compiler = KnowledgeCompiler(store=store)
    document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)

    await compiler.compile_document(document)
    first_evidence = len(store.evidence)
    first_facts = len(store.facts)

    await compiler.compile_document(document)

    assert len(store.compilations) == 2
    assert all(item["status"] == "completed" for item in store.compilations)
    assert len(store.facts) == first_facts, "la recompilación duplicó hechos"
    assert len(store.evidence) > first_evidence, "la recompilación no re-enlazó evidencia"


@pytest.mark.asyncio
async def test_compilador_emite_eventos_semanticos_reales() -> None:
    """La UI aprende de eventos reales: descubierto vs reforzado se distinguen."""
    store = FakeCompilerStore()
    compiler = KnowledgeCompiler(store=store)
    observed: list[tuple[str, dict]] = []

    async def observer(event_type: str, payload: dict) -> None:
        observed.append((event_type, payload))

    first = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)
    result = await compiler.compile_document(first, observer=observer)
    types = [event for event, _ in observed]
    assert "ENTITY_DISCOVERED" in types
    assert "FACT_DISCOVERED" in types
    assert "EVIDENCE_LINKED" in types
    assert "KNOWLEDGE_OBJECT_CREATED" in types
    assert result.persisted["entities_new"] > 0
    assert result.persisted["facts_new"] > 0

    observed.clear()
    second = parse_pdf(record2_pdf(), external_id="record2.pdf", source_id=SOURCE_B)
    await compiler.compile_document(second, observer=observer)
    types = [event for event, _ in observed]
    assert "ENTITY_MATCHED" in types, "la segunda fuente debe reconocer lo ya sabido"
    assert "FACT_REINFORCED" in types, "los hechos compartidos deben reforzarse"


@pytest.mark.asyncio
async def test_compilador_no_rompe_si_el_store_falla() -> None:
    class BrokenStore(FakeCompilerStore):
        async def existing_aliases(self, organization_id):
            raise RuntimeError("base caída")

    compiler = KnowledgeCompiler(store=BrokenStore())
    document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)
    result = await compiler.compile_document(document)
    # El conocimiento determinista se calculó igual; el fallo queda registrado.
    assert result.facts, "el build determinista no debe depender del store"
    assert result.persisted.get("status") == "failed"


# ---------------------------------------------------------------------------
# 5. Eventos de sistema (C8): el pipeline mapea señales reales, best-effort
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_compilador_emite_eventos_de_sistema_desde_el_pipeline_real() -> None:
    """Entidad nueva y regla nueva viajan al emisor de sistema con sus ids."""
    store = FakeCompilerStore()
    emitter = FakeSystemEmitter()
    compiler = KnowledgeCompiler(store=store)
    document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)

    await compiler.compile_document(document, system_emitter=emitter)

    new_entities = emitter.of_type(KnowledgeEventType.NEW_ENTITY)
    assert new_entities, "la entidad nueva debe emitirse al sistema"
    assert all(event.object_id is not None for event in new_entities)
    assert all(event.confidence is not None for event in new_entities)
    assert all(event.document_id == document.id for event in new_entities)
    assert all(event.source_id == SOURCE_A for event in new_entities)
    assert all(event.requires_review is False for event in new_entities)
    assert all(event.event_name == "knowledge.new_entity" for event in new_entities)

    new_rules = emitter.of_type(KnowledgeEventType.NEW_RULE)
    assert new_rules, "la regla nueva debe emitirse al sistema"
    assert all(event.rule_key for event in new_rules)
    assert all(event.object_id is not None for event in new_rules)
    assert all(event.organization_id == ORG for event in emitter.events)
    assert emitter.of_type(KnowledgeEventType.RULE_CHANGED) == []


@pytest.mark.asyncio
async def test_compilador_emite_rule_changed_si_el_sujeto_ya_tenia_otra_regla() -> None:
    """Mismo subject con distinta rule_key: la regla previa queda a revisión."""
    store = FakeCompilerStore()
    emitter = FakeSystemEmitter()
    first_rule = RuleCandidate(
        subject="Record 4",
        statement="Record 4 must renumber on reissue.",
        rule_key="key-v1",
    )
    changed_rule = RuleCandidate(
        subject="Record 4",
        statement="Record 4 must renumber within 24 hours.",
        rule_key="key-v2",
    )
    document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)

    first = PreparedCompiler(_prepared_result(rules=[first_rule]), store=store)
    await first.compile_document(document, system_emitter=emitter)
    assert emitter.of_type(KnowledgeEventType.NEW_RULE), "la primera regla es nueva"
    assert emitter.of_type(KnowledgeEventType.RULE_CHANGED) == []
    emitter.events.clear()

    second = PreparedCompiler(_prepared_result(rules=[changed_rule]), store=store)
    await second.compile_document(document, system_emitter=emitter)

    changed = emitter.of_type(KnowledgeEventType.RULE_CHANGED)
    assert len(changed) == 1, "el cambio de regla debe emitirse una sola vez"
    event = changed[0]
    assert event.payload == {
        "previous_rule_key": "key-v1",
        "rule_key": "key-v2",
        "subject": "Record 4",
    }
    assert event.rule_key == "key-v2"
    assert event.requires_review is True
    assert emitter.of_type(KnowledgeEventType.NEW_RULE) == []


@pytest.mark.asyncio
async def test_compilador_emite_conflicto_de_sistema_con_tipo_e_ids() -> None:
    """El conflicto real del pipeline viaja con tipo, ids y revisión humana."""
    store = FakeCompilerStore()
    emitter = FakeSystemEmitter()
    evidence_id = uuid4()
    conflict = ConflictCandidate(
        subject="Byte 105",
        predicate="has_length",
        value_a="3",
        value_b="5",
        conflict_type=ConflictType.SOURCE_CONFLICT.value,
        confidence=0.8,
        reason="dos fuentes independientes declaran longitudes distintas",
        evidence_ids=[evidence_id],
    )
    compiler = PreparedCompiler(_prepared_result(conflicts=[conflict]), store=store)
    document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)

    await compiler.compile_document(document, system_emitter=emitter)

    events = emitter.of_type(KnowledgeEventType.CONFLICT_DETECTED)
    assert len(events) == 1
    event = events[0]
    assert event.payload["conflict_type"] == ConflictType.SOURCE_CONFLICT.value
    assert event.payload["evidence_ids"] == [str(evidence_id)]
    assert event.confidence == 0.8
    assert event.requires_review is True


@pytest.mark.asyncio
async def test_compilador_sin_emisor_de_sistema_no_cambia_el_flujo() -> None:
    """Sin emisor no hay lookups nuevos ni eventos: comportamiento intacto."""

    class RecordingStore(FakeCompilerStore):
        def __init__(self) -> None:
            super().__init__()
            self.rule_key_lookups = 0

        async def existing_rule_keys(self, organization_id):
            self.rule_key_lookups += 1
            return await super().existing_rule_keys(organization_id)

    store = RecordingStore()
    compiler = KnowledgeCompiler(store=store)
    document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)

    result = await compiler.compile_document(document)

    assert result.persisted["status"] == "completed"
    assert store.rules, "las reglas se siguen persistiendo sin emisor"
    assert store.rule_key_lookups == 0


@pytest.mark.asyncio
async def test_fallos_del_emisor_y_del_indice_de_reglas_no_rompen_el_pipeline() -> None:
    """Emisión e índice de reglas son best-effort: nunca tumban la compilación."""

    class BrokenEmitter:
        async def emit(self, event) -> None:
            raise RuntimeError("bus caído")

    class BrokenRuleIndex(FakeCompilerStore):
        async def existing_rule_keys(self, organization_id):
            raise RuntimeError("índice caído")

    document = parse_pdf(record4_pdf(), external_id="record4.pdf", source_id=SOURCE_A)

    broken_emitter = await KnowledgeCompiler(
        store=BrokenRuleIndex()
    ).compile_document(document, system_emitter=BrokenEmitter())
    assert broken_emitter.persisted["status"] == "completed"

    emitter = FakeSystemEmitter()
    broken_index = await KnowledgeCompiler(store=BrokenRuleIndex()).compile_document(
        document, system_emitter=emitter
    )
    assert broken_index.persisted["status"] == "completed"
    assert emitter.of_type(KnowledgeEventType.NEW_RULE), (
        "sin índice de reglas, toda regla creada es nueva"
    )
