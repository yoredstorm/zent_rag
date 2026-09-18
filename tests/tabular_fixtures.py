# =============================================================================
# Fixtures tabulares para tests (xlsx/csv en memoria, sin binarios en repo)
# =============================================================================
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font


def _save(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def atpco_workbook_bytes() -> bytes:
    """Caso §43: título + subtítulo + blank + header en fila 4."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Record2"
    sheet["A1"] = "ATPCO RECORD 2 RULES"
    sheet["A1"].font = Font(bold=True)
    sheet["A2"] = "Field layout specification"
    headers = ["Field Name", "Start Position", "End Position", "Length", "Description"]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=column, value=header)
        cell.font = Font(bold=True)
    rows = [
        ["Carrier Code", 28, 29, 2, "Identifies carrier"],
        ["Tariff Number", 30, 32, 3, "Identifies tariff"],
        ["Rule Number", 33, 35, 3, "Identifies rule"],
        ["Fare Class", 36, 43, 8, "Fare class code"],
    ]
    for row_index, row in enumerate(rows, start=5):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)
    return _save(workbook)


def multi_sheet_workbook_bytes() -> bytes:
    workbook = Workbook()
    rules = workbook.active
    rules.title = "Rules"
    for column, header in enumerate(["Code", "Description", "Start", "Length"], start=1):
        rules.cell(row=1, column=column, value=header)
    for row_index, row in enumerate(
        [
            ["R2", "Record 2 rules", 28, 2],
            ["R3", "Record 3 rules", 40, 4],
        ],
        start=2,
    ):
        for column, value in enumerate(row, start=1):
            rules.cell(row=row_index, column=column, value=value)

    carriers = workbook.create_sheet("Carriers")
    for column, header in enumerate(["Carrier Code", "Name"], start=1):
        carriers.cell(row=1, column=column, value=header)
    for row_index, row in enumerate(
        [["AA", "American Airlines"], ["AM", "Aeromexico"]], start=2
    ):
        for column, value in enumerate(row, start=1):
            carriers.cell(row=row_index, column=column, value=value)
    return _save(workbook)


def header_not_first_row_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "Reporte de posiciones"
    sheet["A2"] = "Generado 2026-01-15"
    sheet["A3"] = "Área: Operaciones"
    for column, header in enumerate(["Field Name", "Position", "Length"], start=1):
        sheet.cell(row=5, column=column, value=header)
    for row_index, row in enumerate(
        [["Carrier Code", 28, 2], ["Tariff", 30, 3]], start=6
    ):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)
    return _save(workbook)


def multi_header_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Record2"
    sheet["A1"] = "ATPCO RECORD 2 RULES"
    sheet.merge_cells("B3:C3")
    sheet["B3"] = "Position"
    for column, header in enumerate(
        ["Field Name", "Start", "End", "Length"], start=1
    ):
        sheet.cell(row=4, column=column, value=header)
    for row_index, row in enumerate(
        [["Carrier Code", 28, 29, 2], ["Tariff Number", 30, 32, 3]], start=5
    ):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)
    return _save(workbook)


def merged_headers_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Merged"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "FIELD DEFINITION"
    sheet["A3"] = "Field Name"
    sheet["B3"] = "Position"
    sheet["C3"] = "Length"
    for row_index, row in enumerate(
        [["Carrier Code", 28, 2], ["Tariff", 30, 3]], start=4
    ):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)
    return _save(workbook)


def multiple_tables_sheet_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Multi"
    sheet["A1"] = "TABLE A: Field definitions"
    for column, header in enumerate(["Field Name", "Position"], start=1):
        sheet.cell(row=2, column=column, value=header)
    for row_index, row in enumerate([["Carrier", 28], ["Tariff", 30]], start=3):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)
    sheet["A6"] = "TABLE B: Category mappings"
    for column, header in enumerate(["Category", "Meaning"], start=1):
        sheet.cell(row=7, column=column, value=header)
    for row_index, row in enumerate([["CAT10", "Pricing"], ["CAT14", "Rules"]], start=8):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)
    return _save(workbook)


def dates_and_formulas_workbook_bytes() -> bytes:
    import datetime as dt

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fechas"
    for column, header in enumerate(["Item", "Valid From", "Price", "Total"], start=1):
        sheet.cell(row=1, column=column, value=header)
    sheet["A2"] = "Basic"
    sheet["B2"] = dt.date(2026, 1, 15)
    sheet["C2"] = 10.5
    sheet["D2"] = "=C2*2"
    sheet["A3"] = "Premium"
    sheet["B3"] = "15/02/2026"
    sheet["C3"] = 20
    sheet["D3"] = "=C3*2"
    return _save(workbook)


def codes_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Codes"
    for column, header in enumerate(["Category", "Record", "Value", "Fare Basis"], start=1):
        sheet.cell(row=1, column=column, value=header)
    rows = [
        ["CAT10", "R2", "978", "YQYR"],
        ["CAT 14", "Record2", "R&&&&&E&", "FCLASS"],
        ["TARNO", "Loc1", "Loc2", "R2"],
    ]
    for row_index, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)
    return _save(workbook)


def hidden_rows_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hidden"
    for column, header in enumerate(["Code", "Value"], start=1):
        sheet.cell(row=1, column=column, value=header)
    sheet["A2"], sheet["B2"] = "A", 1
    sheet["A3"], sheet["B3"] = "B", 2
    sheet["A4"], sheet["B4"] = "C", 3
    sheet.row_dimensions[3].hidden = True
    return _save(workbook)


def excel_table_workbook_bytes() -> bytes:
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Formal"
    sheet["A1"] = "Zone"
    sheet["B1"] = "Code"
    for row_index, row in enumerate([["North", "N1"], ["South", "S1"]], start=2):
        sheet.cell(row=row_index, column=1, value=row[0])
        sheet.cell(row=row_index, column=2, value=row[1])
    table = Table(displayName="ZoneTable", ref="A1:B3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    sheet.add_table(table)
    return _save(workbook)


def duplicate_headers_workbook_bytes() -> bytes:
    """Header real con columna homónima ("Data Row" dos veces) y datos densos.

    Replica el caso ATPCO Attributes: el header NO es único y la fila de datos
    mezcla números, identificadores (1_100) y texto. La detección no debe
    rechazar la tabla por headers duplicados.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ATPCO_Attributes"
    headers = [
        "Data Row",
        "Std #",
        "Item Row No.",
        "Standard Name",
        "Location",
        "Field Name",
        "Status",
        "Data Row",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=column, value=header)
    rows = [
        ["24", "101", "1", "Date Table 157", "1_100", "Date Table 157", "Current", "24"],
        ["25", "101", "2", "Date Table 157", "1", "Record Type", "Current", "25"],
        ["26", "101", "3", "Date Table 157", "2", "Action", "Current", "26"],
    ]
    for row_index, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)
    return _save(workbook)


def large_workbook_bytes(rows: int = 2_000) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Large")
    sheet.append(["Field Name", "Position", "Length"])
    for index in range(rows):
        sheet.append([f"Field {index:05d}", index, index % 9 + 1])
    return _save(workbook)


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------


def simple_csv_bytes() -> bytes:
    return (
        "Field Name,Start Position,Length\n"
        "Carrier Code,28,2\n"
        "Tariff Number,30,3\n"
    ).encode("utf-8")


def semicolon_csv_bytes() -> bytes:
    return (
        "Field Name;Start Position;Length\nCarrier Code;28;2\nTariff Number;30;3\n"
    ).encode("utf-8")


def utf8_bom_csv_bytes() -> bytes:
    return (
        "\ufeffField Name,Start Position,Length\nCarrier Code,28,2\nTariff Number,30,3\n"
    ).encode("utf-8")


def latin1_csv_bytes() -> bytes:
    text = "Código;Descripción;Precio\nCAT10;Precio básico;10,50\n"
    return text.encode("latin-1")


def quoted_csv_bytes() -> bytes:
    return (
        'Name,Description,Value\n'
        '"Carrier, Code","Identifies carrier, primary",28\n'
    ).encode("utf-8")


def ragged_csv_bytes() -> bytes:
    return (
        "A,B,C\n1,2,3\n4,5\n6,7,8,9\n"
    ).encode("utf-8")
