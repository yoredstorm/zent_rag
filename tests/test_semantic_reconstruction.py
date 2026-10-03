# =============================================================================
# Semantic Reconstruction Layer — suite multiformato
# =============================================================================
# Contrato (spec §26/§27): estructura, continuidad, provenance, semantic units
# y gate se preservan en PDF, DOCX, TXT, Markdown, HTML, Excel, CSV, JSON, XML,
# base de datos y APIs. Un fragmento sin evidencia NUNCA es conocimiento.
# =============================================================================
from __future__ import annotations

import io
import json
from uuid import uuid4

import pytest

from src.core.domain.knowledge_v2 import (
    BoundingBox,
    DocumentPage,
    DocumentTable,
    StructuredBlock,
    StructuredBlockKind,
    StructuredDocument,
)
from src.knowledge.compiler import KnowledgeCompiler
from src.knowledge.reconstruction import (
    ReconstructionStatus,
    escalate_reconstruction,
    reconstruct_document,
    reconstruction_quality_issues,
    registered_kinds,
    tech_view,
)
from src.knowledge.reconstruction.contracts import (
    ElementKind,
    RawElement,
    SourceProvenance,
    element_uuid,
)
from src.knowledge.reconstruction.fragments import evaluate_fragment
from src.knowledge.reconstruction.quality_gate import gate_element
from src.knowledge.structure import CsvParser, DocxParser, TextParser, XlsxParser
from src.knowledge.understanding.engine import apply_understanding

ORG = uuid4()


def _block(
    text: str,
    order: int,
    page: int = 1,
    *,
    kind: StructuredBlockKind = StructuredBlockKind.PARAGRAPH,
    y0: float = 100.0,
    x0: float = 72.0,
    x1: float = 420.0,
    metadata: dict | None = None,
) -> StructuredBlock:
    return StructuredBlock(
        kind=kind,
        text=text,
        order=order,
        page=page,
        bbox=BoundingBox(page=page, x0=x0, y0=y0, x1=x1, y1=y0 + 12),
        metadata=dict(metadata or {}),
    )


def _document(
    blocks: tuple[StructuredBlock, ...] | list[StructuredBlock],
    *,
    title: str = "Source",
    external_id: str = "source.pdf",
    tables: tuple[DocumentTable, ...] = (),
    pages: tuple[DocumentPage, ...] = (),
    metadata: dict | None = None,
) -> StructuredDocument:
    document_id = uuid4()
    rebound_tables = tuple(
        DocumentTable(
            id=table.id,
            document_id=document_id,
            organization_id=ORG,
            source_id=SOURCE_ID,
            headers=table.headers,
            rows=table.rows,
            page=table.page,
            caption=table.caption,
            metadata=dict(table.metadata),
        )
        for table in tables
    )
    return StructuredDocument(
        id=document_id,
        organization_id=ORG,
        source_id=SOURCE_ID,
        external_id=external_id,
        title=title,
        content_hash="pending",
        blocks=tuple(blocks),
        tables=rebound_tables,
        pages=tuple(
            DocumentPage(
                id=page.id,
                document_id=document_id,
                organization_id=ORG,
                source_id=SOURCE_ID,
                page_number=page.page_number,
                text=page.text,
                block_ids=page.block_ids,
                metadata=dict(page.metadata),
            )
            for page in pages
        ),
        metadata=dict(metadata or {}),
    )


SOURCE_ID = uuid4()


def _payload(document: StructuredDocument) -> dict:
    return document.metadata["semantic_reconstruction"]


UPDATE_EXPECTED = {
    ReconstructionStatus.VALID.value,
    ReconstructionStatus.RECONSTRUCTED.value,
}


# ---------------------------------------------------------------------------
# 1 · PDF: continuidad, palabra partida, columnas y página
# ---------------------------------------------------------------------------


def test_pdf_word_fragment_reconstructed_with_structural_evidence() -> None:
    document = _document(
        [
            _block("FOR RECORD 2 - CATEG", 0, y0=100),
            _block("ORY CONTROL continues here", 1, y0=112),
        ],
        title="DATA APPLICATION FOR RECORD 2 - CATEGORY CONTROL",
    )
    outcome = reconstruct_document(document)
    blocks = outcome.document.blocks

    assert blocks[0].text == "FOR RECORD 2 - CATEGORY CONTROL continues here"
    assert blocks[0].metadata["reconstruction_status"] == ReconstructionStatus.RECONSTRUCTED.value
    assert blocks[1].metadata["superseded"] is True
    assert blocks[1].metadata["index_semantic"] is False
    assert blocks[1].metadata["superseded_into"] == str(blocks[0].id)
    stats = outcome.report["stats"]
    assert stats["merged_continuations"] == 1
    assert stats["units_reconstructed"] == 1
    # Provenance viva: el bloque absorbido sigue existiendo en el payload.
    units = _payload(outcome.document)["units"]
    assert any(
        unit["provenance"].get("block_id") == str(blocks[1].id)
        and unit["status"] == ReconstructionStatus.DUPLICATE.value
        for unit in units
    )


def test_word_fragment_without_evidence_is_never_merged() -> None:
    document = _document([_block("ZXQ", 0), _block("WVB", 1)], title="t")
    outcome = reconstruct_document(document)
    blocks = outcome.document.blocks

    assert blocks[0].text == "ZXQ"
    assert blocks[1].metadata.get("superseded") is not True
    assert blocks[0].metadata["index_semantic"] is False
    assert blocks[1].metadata["index_semantic"] is False
    pending = _payload(outcome.document)["pending_decisions"]
    assert len(pending) == 1
    assert pending[0]["kind"] in {"WORD_FRAGMENT", "DEHYPHENATION"}


def test_isolated_fragment_never_becomes_canonical_knowledge() -> None:
    document = _document([_block("Y CONT", 0, y0=160)], title="CATEGORY CONTROL MANUAL")
    document = apply_understanding(document, filename="manual.pdf")
    payload = _payload(document)

    status = next(unit["status"] for unit in payload["units"] if unit["label"] == "Y CONT")
    assert status in {
        ReconstructionStatus.INCOMPLETE.value,
        ReconstructionStatus.REQUIRES_REPAIR.value,
        ReconstructionStatus.REJECTED.value,
        ReconstructionStatus.LOW_QUALITY.value,
    }
    block = next(block for block in document.blocks if block.text == "Y CONT")
    assert block.metadata["index_semantic"] is False

    result = KnowledgeCompiler.build(document)
    assert not any(entity.name == "Y CONT" for entity in result.entities)
    assert any(
        issue.kind == "INGESTION_QUALITY" for issue in result.quality_issues
    )


def test_multicolumn_reading_order_does_not_merge_columns() -> None:
    document = _document(
        [
            _block("The fare component amount", 0, x0=72, x1=260),
            _block("applies to all carriers", 1, x0=320, x1=520, y0=112),
        ],
        title="Fare rules",
    )
    outcome = reconstruct_document(document)

    assert outcome.document.blocks[0].text == "The fare component amount"
    assert outcome.document.blocks[1].text == "applies to all carriers"
    assert outcome.report["stats"]["merged_continuations"] == 0


def test_dehyphenation_joins_word_across_lines() -> None:
    document = _document([_block("inter-", 0), _block("national carriers apply", 1)], title="t")
    outcome = reconstruct_document(document)

    assert outcome.document.blocks[0].text == "international carriers apply"
    continuation = outcome.representation.continuations[0]
    assert continuation.kind == "DEHYPHENATION"


def test_page_boundary_only_merges_with_word_evidence() -> None:
    without_evidence = _document(
        [_block("The rule applies", 0, page=1), _block("ALL CARRIERS", 1, page=2)],
        title="Rules",
        pages=(
            DocumentPage(
                id=uuid4(), document_id=uuid4(), organization_id=ORG, page_number=1, text="The rule applies"
            ),
        ),
    )
    outcome = reconstruct_document(without_evidence)
    assert outcome.report["stats"]["merged_continuations"] == 0

    with_evidence = _document(
        [
            _block("RECORD 2 - CATEG", 0, page=1),
            _block("ORY CONTROL", 1, page=2),
        ],
        title="DATA APPLICATION FOR RECORD 2 - CATEGORY CONTROL",
    )
    outcome = reconstruct_document(with_evidence)
    assert outcome.document.blocks[0].text.endswith("CATEGORY CONTROL")
    assert outcome.report["stats"]["merged_continuations"] == 1


def test_repeated_header_is_structural_artifact_not_knowledge() -> None:
    blocks = [
        _block("CARRIER MANUAL", 0, page=1, y0=40),
        _block("Body text of page one.", 1, page=1, y0=200),
        _block("CARRIER MANUAL", 2, page=2, y0=40),
        _block("Body text of page two.", 3, page=2, y0=200),
    ]
    document = apply_understanding(_document(blocks, title="Manual"), filename="manual.pdf")
    payload = _payload(document)
    statuses = {
        unit["label"]: unit["status"]
        for unit in payload["units"]
        if unit["label"] == "CARRIER MANUAL"
    }
    assert set(statuses.values()) == {ReconstructionStatus.STRUCTURAL_ARTIFACT.value}
    assert all(
        block.metadata.get("index_semantic") is False
        for block in document.blocks
        if block.text == "CARRIER MANUAL"
    )


# ---------------------------------------------------------------------------
# 2 · Excel / CSV: estructura tabular, no filas a texto
# ---------------------------------------------------------------------------


def _parse_xlsx(data: bytes, external_id: str = "book.xlsx") -> StructuredDocument:
    document = XlsxParser().parse(
        data,
        organization_id=ORG,
        external_id=external_id,
        source_id=SOURCE_ID,
        source_name=external_id,
    )
    return apply_understanding(document, filename=external_id)


def _parse_csv(data: bytes, external_id: str = "data.csv") -> StructuredDocument:
    document = CsvParser().parse(
        data,
        organization_id=ORG,
        external_id=external_id,
        source_id=SOURCE_ID,
        source_name=external_id,
    )
    return apply_understanding(document, filename=external_id)


def test_excel_multi_sheet_preserves_tables_and_records() -> None:
    from tests.tabular_fixtures import multi_sheet_workbook_bytes

    document = _parse_xlsx(multi_sheet_workbook_bytes(), "multi.xlsx")
    payload = _payload(document)

    assert payload["source_kind"] == "spreadsheet"
    assert payload["stats"]["sheets"] == 2
    tables = payload["tables"]
    assert len(tables) == 2
    headers = {tuple(table["headers"]) for table in tables}
    assert ("Code", "Description", "Start", "Length") in headers
    assert ("Carrier Code", "Name") in headers
    carriers = next(table for table in tables if table["headers"] == ["Carrier Code", "Name"])
    assert carriers["rows"] == [["AA", "American Airlines"], ["AM", "Aeromexico"]]
    field_names = {field["name"] for field in carriers["fields"]}
    assert {"Carrier Code", "Name"} <= field_names


def test_csv_preserves_headers_types_and_provenance() -> None:
    data = (
        "Carrier Code,Category,Byte,Meaning\n"
        "AA,Cat 31,105,Currency code\n"
        "UA,Cat 31,106,Fare basis\n"
    ).encode("utf-8")
    document = _parse_csv(data, "carriers.csv")
    payload = _payload(document)

    assert payload["source_kind"] == "csv"
    table = payload["tables"][0]
    assert table["headers"] == ["Carrier Code", "Category", "Byte", "Meaning"]
    assert table["rows"][0] == ["AA", "Cat 31", "105", "Currency code"]
    assert table["provenance"]["sheet"]
    byte_field = next(field for field in table["fields"] if field["name"] == "Byte")
    assert byte_field["inferred_type"] in {"int", "integer", "float", "decimal", "unknown"}
    assert payload["stats"]["schemas_inferred"] >= 1


def test_excel_formulas_survive_reconstruction() -> None:
    from tests.tabular_fixtures import dates_and_formulas_workbook_bytes

    document = _parse_xlsx(dates_and_formulas_workbook_bytes(), "dates.xlsx")
    payload = _payload(document)
    table = payload["tables"][0]

    assert any(formula["formula"] == "=C2*2" for formula in table["formulas"])
    assert table["headers"][0] == "Item"


# ---------------------------------------------------------------------------
# 3 · JSON / XML: jerarquía, rutas y nada aplanado
# ---------------------------------------------------------------------------


def test_json_preserves_nested_hierarchy_and_paths() -> None:
    raw = json.dumps(
        {
            "carrier": {"code": "AM", "name": "Aeromexico"},
            "flights": [{"num": 1}, {"num": 2}],
        }
    )
    document = _document(
        [], title="Carrier payload", external_id="payload", metadata={"source_kind": "json"}
    )
    outcome = reconstruct_document(document, raw_text=raw)
    payload = _payload(outcome.document)

    paths = {
        unit["attributes"].get("json_path") for unit in payload["units"] if unit["kind"] == "property"
    }
    assert "$.carrier.code" in paths
    assert "$.flights[0].num" in paths
    nodes = {node["id"]: node for node in payload["nodes"]}
    child = next(node for node in payload["nodes"] if node["text"] == "carrier")
    parent = nodes.get(child["parent_id"])
    assert parent is not None and parent["text"] == "$"


def test_xml_preserves_xpath_and_attributes() -> None:
    raw = '<carriers><carrier code="AM"><name>Aeromexico</name></carrier></carriers>'
    document = _document(
        [], title="Carrier xml", external_id="payload", metadata={"source_kind": "xml"}
    )
    outcome = reconstruct_document(document, raw_text=raw)
    payload = _payload(outcome.document)

    xpaths = {unit["attributes"].get("xpath") for unit in payload["units"]}
    assert "/carriers[1]/carrier[1]/name[1]" in xpaths
    carrier = next(unit for unit in payload["units"] if unit["label"] == "carrier")
    assert carrier["attributes"]["attributes"] == {"code": "AM"}


# ---------------------------------------------------------------------------
# 4 · Base de datos y API
# ---------------------------------------------------------------------------


def test_database_schema_preserves_keys_relationships_and_types() -> None:
    document = _document(
        [],
        title="DWH",
        external_id="dwh",
        metadata={
            "database_schema": {
                "database": "dwh",
                "schemas": [
                    {
                        "name": "public",
                        "tables": [
                            {
                                "name": "carriers",
                                "description": "Master carrier table",
                                "row_count": 1200,
                                "columns": [
                                    {
                                        "name": "cxrcd",
                                        "data_type": "char(2)",
                                        "primary_key": True,
                                        "nullable": False,
                                    },
                                    {
                                        "name": "cat_id",
                                        "data_type": "integer",
                                        "foreign_key": True,
                                        "references": {"table": "categories", "column": "id"},
                                    },
                                ],
                            }
                        ],
                    }
                ],
            }
        },
    )
    outcome = reconstruct_document(document)
    payload = _payload(outcome.document)

    assert payload["source_kind"] == "database"
    code = next(unit for unit in payload["units"] if unit["label"] == "cxrcd")
    assert code["attributes"]["primary_key"] is True
    assert code["attributes"]["data_type"] == "char(2)"
    category = next(unit for unit in payload["units"] if unit["label"] == "cat_id")
    assert category["attributes"]["references_table"] == "categories"
    table = payload["tables"][0]
    assert table["metadata"]["primary_key"] == ["cxrcd"]
    assert table["metadata"]["foreign_keys"][0]["references_table"] == "categories"


def test_openapi_preserves_endpoints_parameters_and_responses() -> None:
    spec = {
        "openapi": "3.0.0",
        "info": {"title": "Carrier API", "version": "1.2"},
        "paths": {
            "/carriers/{code}": {
                "get": {
                    "operationId": "getCarrier",
                    "parameters": [
                        {
                            "name": "code",
                            "in": "path",
                            "required": True,
                            "schema": {"type": "string"},
                        }
                    ],
                    "responses": {
                        "200": {
                            "description": "OK",
                            "content": {
                                "application/json": {
                                    "schema": {"type": "object", "properties": {"code": {"type": "string"}}}
                                }
                            },
                        }
                    },
                }
            }
        },
    }
    document = _document(
        [], title="Carrier API", external_id="openapi.json", metadata={"source_kind": "api"}
    )
    outcome = reconstruct_document(document, raw_text=json.dumps(spec))
    payload = _payload(outcome.document)

    assert payload["source_kind"] == "api"
    endpoint = next(unit for unit in payload["units"] if unit["kind"] == "endpoint")
    assert endpoint["attributes"]["operation_id"] == "getCarrier"
    assert endpoint["attributes"]["responses"]["200"]["description"] == "OK"
    parameter = next(unit for unit in payload["units"] if unit["kind"] == "parameter")
    assert parameter["label"] == "code"
    response = next(unit for unit in payload["units"] if unit["kind"] == "response")
    assert response["attributes"]["status_code"] == "200"


# ---------------------------------------------------------------------------
# 5 · DOCX / Markdown / texto: jerarquía textual
# ---------------------------------------------------------------------------


def test_docx_table_is_preserved_not_flattened() -> None:
    pytest.importorskip("docx")
    from docx import Document as DocxDocument

    buffer = io.BytesIO()
    doc = DocxDocument()
    doc.add_heading("Manual Operaciones", level=0)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Clave"
    table.cell(0, 1).text = "Valor"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "5%"
    doc.save(buffer)

    parsed = DocxParser().parse(
        buffer.getvalue(),
        organization_id=ORG,
        external_id="manual.docx",
        source_id=SOURCE_ID,
        source_name="manual.docx",
    )
    understood = apply_understanding(parsed, filename="manual.docx")
    payload = _payload(understood)

    assert payload["source_kind"] == "docx"
    assert payload["tables"][0]["headers"] == ["Clave", "Valor"]
    assert payload["tables"][0]["rows"] == [["A", "5%"]]


def test_markdown_hierarchy_is_valid_semantics() -> None:
    markdown = b"# Title\n\nIntro paragraph.\n\n## Section\n\n- item one\n- item two\n"
    parsed = TextParser().parse(
        markdown,
        organization_id=ORG,
        external_id="doc.md",
        source_id=SOURCE_ID,
        source_name="doc.md",
    )
    outcome = reconstruct_document(parsed)
    payload = _payload(outcome.document)

    assert payload["source_kind"] == "markdown"
    statuses = {unit["status"] for unit in payload["units"]}
    assert statuses <= UPDATE_EXPECTED
    labels = {unit["label"] for unit in payload["units"]}
    assert {"Title", "Section"} <= labels


# ---------------------------------------------------------------------------
# 6 · Gate, fragmentos y classifications
# ---------------------------------------------------------------------------


def test_quality_gate_classifies_and_blocks_artifacts() -> None:
    def element(text: str, kind: str = ElementKind.PARAGRAPH.value, **attrs) -> RawElement:
        return RawElement(
            id=element_uuid(uuid4(), 0, text),
            kind=kind,
            text=text,
            order=0,
            provenance=SourceProvenance(excerpt=text),
            attributes=attrs,
        )

    valid = gate_element(
        evaluate_fragment(element("This is a complete sentence."), references=("Intro. This is a complete sentence.",))
    )
    assert valid.status == ReconstructionStatus.VALID.value
    assert valid.knowledge and valid.index_semantic

    chrome = gate_element(evaluate_fragment(element("PAGE 4", kind=ElementKind.HEADER.value)))
    assert chrome.status == ReconstructionStatus.STRUCTURAL_ARTIFACT.value
    assert not chrome.knowledge and not chrome.index_semantic

    fragment = gate_element(
        evaluate_fragment(element("Y CONT"), references=("CATEGORY CONTROL MANUAL",))
    )
    assert fragment.status in {
        ReconstructionStatus.INCOMPLETE.value,
        ReconstructionStatus.REQUIRES_REPAIR.value,
        ReconstructionStatus.LOW_QUALITY.value,
        ReconstructionStatus.REJECTED.value,
    }
    assert not fragment.knowledge and not fragment.index_semantic


def test_reconstruction_quality_issues_are_ingestion_quality() -> None:
    document = _document([_block("Y CONT", 0)], title="CATEGORY CONTROL MANUAL")
    outcome = reconstruct_document(document)

    issues = reconstruction_quality_issues(outcome.document)
    assert issues
    assert {issue["kind"] for issue in issues} == {"INGESTION_QUALITY"}
    assert all(
        issue["detail"]["classification"] not in UPDATE_EXPECTED for issue in issues
    )


# ---------------------------------------------------------------------------
# 7 · LLM: solo ambigüedad real, output estructurado, sin chain-of-thought
# ---------------------------------------------------------------------------


class _FakeProvider:
    name = "fake"

    def __init__(self, response: dict) -> None:
        self.response = response

    async def reconstruct(self, request: dict) -> dict:
        return dict(self.response)


@pytest.mark.asyncio
async def test_llm_escalation_applies_only_verified_continuation() -> None:
    document = _document([_block("ZXQ", 0), _block("WVB", 1)], title="t")
    outcome = reconstruct_document(document)
    assert _payload(outcome.document)["pending_decisions"]

    provider = _FakeProvider(
        {
            "classification": "CONTINUATION",
            "semantic_unit": "ZXQWVB",
            "source_elements": ["left", "right"],
            "confidence": 0.93,
            "ambiguity": False,
            "reason": "adjacent cells",
        }
    )
    updated, usage = await escalate_reconstruction(outcome.document, provider)
    blocks = updated.blocks

    assert blocks[0].text == "ZXQWVB"
    assert blocks[1].metadata["index_semantic"] is False
    assert usage.repairs == 1
    payload = _payload(updated)
    assert payload["llm"]["repairs"] == 1
    assert not payload["pending_decisions"]
    assert "chain_of_thought" not in json.dumps(payload).lower()


@pytest.mark.asyncio
async def test_llm_cannot_invent_content() -> None:
    document = _document([_block("ZXQ", 0), _block("WVB", 1)], title="t")
    outcome = reconstruct_document(document)

    provider = _FakeProvider(
        {
            "classification": "CONTINUATION",
            "semantic_unit": "ZXQ INVENTED WVB",
            "confidence": 0.99,
            "ambiguity": False,
        }
    )
    updated, usage = await escalate_reconstruction(outcome.document, provider)

    assert updated.blocks[0].text == "ZXQ"
    assert usage.repairs == 0
    assert updated.blocks[1].metadata["index_semantic"] is False


# ---------------------------------------------------------------------------
# 8 · Convergencia multiformato y tech view
# ---------------------------------------------------------------------------


def test_registry_covers_all_source_kinds() -> None:
    kinds = set(registered_kinds())
    assert {
        "pdf",
        "docx",
        "text",
        "markdown",
        "html",
        "spreadsheet",
        "csv",
        "json",
        "xml",
        "database",
        "api",
        "event",
        "email",
        "conversation",
    } <= kinds


def test_tech_view_reports_reconstruction_metrics() -> None:
    document = _document(
        [
            _block("FOR RECORD 2 - CATEG", 0, y0=100),
            _block("ORY CONTROL continues here", 1, y0=112),
            _block("Y CONT", 2, y0=160),
        ],
        title="DATA APPLICATION FOR RECORD 2 - CATEGORY CONTROL",
    )
    outcome = reconstruct_document(document)
    view = tech_view(outcome.document)

    assert view["raw_blocks"] == 3
    assert view["merged_continuations"] == 1
    assert view["rejected_fragments"] >= 1
    assert view["deterministic"] == 1


def test_compiler_uses_reconstructed_blocks() -> None:
    document = _document(
        [
            _block("FOR RECORD 2 - CATEG", 0, y0=100),
            _block("ORY CONTROL continues here", 1, y0=112),
            _block("Y CONT", 2, y0=160),
        ],
        title="DATA APPLICATION FOR RECORD 2 - CATEGORY CONTROL",
    )
    outcome = reconstruct_document(document)
    result = KnowledgeCompiler.build(outcome.document)

    labels = {unit.label for unit in result.units}
    assert "CATEG" not in labels
    assert "Y CONT" not in labels
    assert not any(entity.name in {"CATEG", "Y CONT"} for entity in result.entities)
