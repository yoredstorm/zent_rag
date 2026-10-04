# =============================================================================
# Semantic Enrichment — tests deterministas
# =============================================================================
# Reglas que se prueban:
#   - el ejemplo CATEGORY 31 -> "Voluntary Changes" + aliases (CAT31, Cat 31)
#   - identificadores verbatim (Byte 105)
#   - acrónimos explícitos e iniciales
#   - temporal (incluye ambigüedad marcada, nunca decidida)
#   - preguntas sintéticas acotadas y deduplicadas
#   - SIN alucinación: todo item tiene provenance a unidades reales
#   - canonical/evidence jamás se mezclan con enrichment
# =============================================================================
from __future__ import annotations

from uuid import uuid4

import pytest

from src.knowledge.enrichment import (
    EnrichmentProfilePack,
    enrich_document,
    register_profile_pack,
)
from src.knowledge.enrichment.contracts import SemanticConcept
from src.knowledge.enrichment.profiling import ProfilePattern
from src.knowledge.enrichment.versioning import ENRICHMENT_VERSION
from src.knowledge.structure.text_parser import TextParser
from src.knowledge.understanding.engine import apply_understanding

MANUAL = """# Category 31 - Voluntary Changes

Category 31 (CAT31) defines voluntary changes for exchange eligibility.

Field: Status Byte
Bytes: 105-105
Format: 1
Description: status of the voluntary change

Byte 105 indicates the status for voluntary changes.

Effective from 2024-01-01 to 2024-12-31. Expiration 01/02/2024.

If the status is A, then the change is accepted.
The passenger must reissue the ticket before departure.

Amount: 150.50 USD and 5% penalty.
"""


def _understood(text: str = MANUAL, external_id: str = "manual.md"):
    parser = TextParser()
    document = parser.parse(
        text.encode("utf-8"),
        organization_id=uuid4(),
        external_id=external_id,
        source_id=uuid4(),
        source_name=external_id,
    )
    return apply_understanding(document, filename=external_id)


@pytest.fixture(scope="module")
def enrichment():
    document = _understood()
    return document, enrich_document(document)


def _all_items(result):
    return list(result.items())


def test_enrichment_version_is_persisted(enrichment) -> None:
    _document, result = enrichment
    assert result.enrichment_version.startswith(ENRICHMENT_VERSION)
    assert result.policy_version
    assert result.derived is True
    assert result.canonical is False
    assert result.statistics.seconds >= 0.0


def test_category_31_gets_aliases(enrichment) -> None:
    _document, result = enrichment
    names = {concept.canonical_name.casefold() for concept in result.concepts}
    assert any("category 31" in name for name in names), names
    concept = next(
        concept for concept in result.concepts if "category 31" in concept.canonical_name.casefold()
    )
    assert concept.semantic_type == "domain_category"
    aliases = {alias.casefold() for alias in concept.aliases}
    assert aliases, "debe haber aliases derivados"
    assert aliases & {"category31", "category 31", "cat 31", "category-31"}, aliases


def test_explicit_acronym_is_linked(enrichment) -> None:
    _document, result = enrichment
    acronyms = {alias.value.casefold(): alias for alias in result.retrieval_aliases}
    assert "cat31" in acronyms, sorted(acronyms)
    assert acronyms["cat31"].target_concept_id


def test_identifiers_are_verbatim(enrichment) -> None:
    _document, result = enrichment
    values = {identifier.value for identifier in result.identifiers}
    assert "Byte 105" in values, values
    byte_identifier = next(item for item in result.identifiers if item.value == "Byte 105")
    assert byte_identifier.identifier_type == "byte"
    assert byte_identifier.source_unit_ids


def test_temporal_detects_dates_and_marks_ambiguity(enrichment) -> None:
    _document, result = enrichment
    values = {item.value for item in result.temporal_qualifiers}
    assert "2024-01-01" in values
    ambiguous = [item for item in result.temporal_qualifiers if item.ambiguous]
    assert any(item.value == "01/02/2024" for item in ambiguous)


def test_possible_rules_are_candidates_not_knowledge(enrichment) -> None:
    _document, result = enrichment
    assert result.possible_rules
    for rule in result.possible_rules:
        assert rule.derived is True
        assert rule.canonical is False
        assert rule.source_unit_ids
        assert rule.rule_key


def test_questions_are_deduplicated_and_bounded(enrichment) -> None:
    _document, result = enrichment
    questions = [item.question for item in result.synthetic_questions]
    assert questions
    assert len(questions) == len({question.casefold() for question in questions})
    assert len(questions) <= 120
    assert any("Category 31" in question for question in questions)
    assert any("Byte 105" in question for question in questions)
    for question in result.synthetic_questions:
        assert question.synthetic is True
        assert question.canonical is False
        assert question.source_unit_ids


def test_no_hallucination_every_item_has_real_provenance(enrichment) -> None:
    document, result = enrichment
    units = set()
    for block in document.blocks:
        units.add(str(block.id))
    for section in document.sections:
        units.add(str(section.id))
    assert result.quality.source_coverage == 1.0
    assert result.quality.prohibited_items == 0
    for item in _all_items(result):
        assert item.derived is True
        assert item.canonical is False
        assert item.source_unit_ids, f"item sin provenance: {item}"
        assert all(unit in units for unit in item.source_unit_ids)


def test_index_fields_are_bounded_and_derived(enrichment) -> None:
    _document, result = enrichment
    fields = result.index_fields(limit=16)
    assert fields["enrichment_derived"] == "true"
    assert fields["enrichment_canonical"] == "false"
    assert len(fields["concept_ids"]) <= 16
    assert len(fields["retrieval_aliases"]) <= 16
    assert len(fields["retrieval_questions"]) <= 16


def test_domain_terms_units(enrichment) -> None:
    _document, result = enrichment
    terms = {term.term.casefold() for term in result.domain_terms}
    assert any("usd" in term for term in terms), terms
    unit_terms = [term for term in result.domain_terms if term.semantic_type.startswith("unit:")]
    assert unit_terms


def test_acronym_noise_is_not_invented() -> None:
    """'Y CONT' no produce 'YC' ni expande a nada: sin evidencia no hay alias."""
    document = _understood(
        "# Notes\n\nY CONT appears in a column header.\n", external_id="noise.md"
    )
    result = enrich_document(document)
    acronyms = {
        alias.value.upper()
        for alias in result.retrieval_aliases
        if alias.kind == "acronym"
    }
    assert "YC" not in acronyms, acronyms
    names = {concept.canonical_name.casefold() for concept in result.concepts}
    assert "category control" not in names
    assert "y cont" not in names or True  # el término crudo puede existir, no la expansión


def test_cat31_does_not_expand_without_source_evidence() -> None:
    """'CAT 31' no se convierte en 'Voluntary Changes' si la fuente no lo dice."""
    document = _understood(
        "# Codes\n\nCAT 31 is an internal code used by the system.\n",
        external_id="cat31.md",
    )
    result = enrich_document(document)
    names = {concept.canonical_name.casefold() for concept in result.concepts}
    aliases = {
        alias.value.casefold() for alias in result.retrieval_aliases
    }
    assert "voluntary changes" not in names
    assert "voluntary changes" not in aliases
    assert "category 31" not in names  # no hay evidencia de que CAT 31 sea Category 31
    for item in result.items():
        assert item.canonical is False
        assert item.derived is True


def test_tabular_enrichment_has_real_provenance() -> None:
    """Conceptos de columnas tabulares deben apuntar a unidades tabulares reales."""
    import io

    from openpyxl import Workbook

    from src.knowledge.tabular.builder import build_tabular_workbook
    from src.knowledge.tabular.xlsx_reader import read_xlsx_grid

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.cell(row=1, column=1, value="FPROC")
    sheet.cell(row=1, column=2, value="CCUST")
    sheet.cell(row=2, column=1, value="20260101")
    sheet.cell(row=2, column=2, value="C001")
    sheet.cell(row=3, column=1, value="20260102")
    sheet.cell(row=3, column=2, value="C002")
    buffer = io.BytesIO()
    workbook.save(buffer)
    grid = read_xlsx_grid(buffer.getvalue(), "prov.xlsx")
    tabular = build_tabular_workbook(
        grid,
        organization_id=uuid4(),
        external_id="prov.xlsx",
    )
    from dataclasses import replace

    from src.knowledge.structure.text_parser import TextParser

    base = TextParser().parse(
        b"# Workbook\n\nplaceholder\n",
        organization_id=tabular.organization_id,
        external_id="prov.xlsx",
        source_name="prov.xlsx",
    )
    document = replace(base, tabular=tabular)
    result = enrich_document(document)
    tabular_concepts = [
        concept
        for concept in result.concepts
        if concept.semantic_type.startswith("column:")
        or concept.canonical_name in ("FPROC", "CCUST")
    ]
    assert tabular_concepts, [concept.canonical_name for concept in result.concepts]
    assert result.quality.source_coverage == 1.0
    for concept in tabular_concepts:
        assert concept.source_unit_ids, concept.canonical_name
    # Sin huérfanos exportados: todo item tiene provenance real.
    assert result.quality.orphan_items == 0


def test_model_items_without_provenance_are_rejected() -> None:
    document = _understood("# Manual\n\nStatus byte definition.\n")
    bogus = SemanticConcept(
        concept_id="bogus",
        canonical_name="Invented Concept",
        semantic_type="concept",
        source_unit_ids=(),
        confidence=0.99,
    )
    canonical_claim = SemanticConcept(
        concept_id="bogus-2",
        canonical_name="Canonical Claim",
        semantic_type="concept",
        source_unit_ids=(str(document.blocks[0].id),),
        confidence=0.99,
        derived=False,
        canonical=True,
    )
    result = enrich_document(document, model_items=[bogus, canonical_claim])
    assert result.quality.prohibited_items >= 1
    assert result.quality.orphan_items == 0
    names = {concept.canonical_name for concept in result.concepts}
    assert "Invented Concept" not in names
    assert "Canonical Claim" not in names


def test_profile_pack_extension_is_not_hardcoded() -> None:
    import re

    pack = EnrichmentProfilePack(
        name="test_pack",
        version="1",
        patterns=(
            ProfilePattern(
                name="room",
                regex=re.compile(r"\bRoom\s+(\d{2,4})\b", re.IGNORECASE),
                semantic_type="location_room",
                confidence=0.9,
                canonical_template="Room {1}",
            ),
        ),
        term_map={"sala": "room"},
        term_types={"folio": "document_number"},
    )
    register_profile_pack(pack)
    document = _understood("# Guest Notes\n\nRoom 415 requires a key.\n", external_id="room.md")
    result = enrich_document(document, packs=[pack])
    semantic_types = {concept.semantic_type for concept in result.concepts}
    assert "location_room" in semantic_types


def test_enrichment_result_payload_has_no_raw_evidence_text() -> None:
    document = _understood()
    result = enrich_document(document)
    payload = result.payload()
    serialized = str(payload)
    # El payload lleva términos derivados, no las oraciones de evidencia.
    raw_sentence = "Category 31 (CAT31) defines voluntary changes for exchange eligibility."
    assert raw_sentence not in serialized
    assert payload["derived"] is True
    assert payload["canonical"] is False
