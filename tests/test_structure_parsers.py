# =============================================================================
# Knowledge V2 — Structured parsers (Phase B: PDF + Markdown/TXT)
# =============================================================================
from __future__ import annotations

import zlib
from uuid import uuid4

import pytest

from src.core.domain.knowledge_v2 import (
    KnowledgeObjectStatus,
    StructuredBlockKind,
)
from src.knowledge.structure import (
    PdfParser,
    StructuredParserError,
    TextParser,
    get_parser,
    supported_extensions,
)


def _tiny_pdf(text_lines: list[str]) -> bytes:
    """PDF mínimo de una página con varias líneas de texto (FlateDecode)."""
    show = "".join(f"BT /F1 12 Tf 72 {y} Td ({t}) Tj ET\n" for t, y in zip(text_lines, range(720, 680, -20)))
    payload = show.encode("latin-1")
    stream = zlib.compress(payload)
    objs = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj",
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj",
        (
            b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj"
        ),
        b"4 0 obj << /Length %d /Filter /FlateDecode >> stream\n" % len(stream)
        + stream
        + b"\nendstream endobj",
        b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj",
    ]
    head = b"%PDF-1.4\n"
    body = b""
    offsets: list[int] = []
    pos = len(head)
    for obj in objs:
        offsets.append(pos)
        body += obj + b"\n"
        pos += len(obj) + 1
    xref_pos = pos
    xref = (
        b"xref\n0 6\n0000000000 65535 f \n"
        + b"".join(f"{off:010d} 00000 n \n".encode() for off in offsets)
    )
    trailer = (
        b"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n"
        + str(xref_pos).encode()
        + b"\n%%EOF\n"
    )
    return head + body + xref + trailer


_MARKDOWN = """# Manual Operaciones

## 5. Compensación

Comisión 5% para contratos vigentes.

### 5.2 Comisiones

| Clave | Valor |
| --- | --- |
| A | 5% |
| B | 7% |

- cláusula uno
- cláusula dos

```python
comision = 0.05
```
"""


def _parse_markdown():
    return TextParser().parse(
        _MARKDOWN.encode("utf-8"),
        organization_id=uuid4(),
        external_id="manual-operaciones.md",
        source_name="manual-operaciones.md",
    )


def test_markdown_parser_builds_blocks_sections_and_tables() -> None:
    doc = _parse_markdown()
    doc.check_consistency()

    kinds = {b.kind for b in doc.blocks}
    assert StructuredBlockKind.HEADING in kinds
    assert StructuredBlockKind.PARAGRAPH in kinds
    assert StructuredBlockKind.LIST in kinds
    assert StructuredBlockKind.TABLE in kinds
    assert StructuredBlockKind.CODE in kinds

    # 3 headings (Manual, 5. Compensación, 5.2 Comisiones)
    headings = [b for b in doc.blocks if b.kind is StructuredBlockKind.HEADING]
    assert len(headings) == 3

    # árbol de secciones: raíz + 2 hijos anidados
    assert doc.section_count == 3
    assert [s.section_path for s in doc.sections] == [
        ("Manual Operaciones",),
        ("Manual Operaciones", "5. Compensación"),
        ("Manual Operaciones", "5. Compensación", "5.2 Comisiones"),
    ]
    root = doc.sections[0]
    child = doc.sections[1]
    grandchild = doc.sections[2]
    assert root.parent_id is None
    assert child.parent_id == root.id
    assert grandchild.parent_id == child.id

    # tabla markdown preservada
    assert doc.table_count == 1
    table = doc.tables[0]
    assert table.headers == ("Clave", "Valor")
    assert table.rows == (("A", "5%"), ("B", "7%"))

    assert doc.pages == ()
    assert doc.provenance.value == "OBSERVED"
    assert doc.status is KnowledgeObjectStatus.OBSERVED


def test_markdown_parser_title_from_first_heading() -> None:
    doc = _parse_markdown()
    assert doc.title == "Manual Operaciones"
    assert doc.mime_type == "text/markdown"


def test_text_plain_parser_produces_paragraph_blocks() -> None:
    doc = TextParser().parse(
        "Hola mundo.\n\nSegundo párrafo con más texto.".encode("utf-8"),
        organization_id=uuid4(),
        external_id="nota.txt",
    )
    doc.check_consistency()
    paragraphs = [b for b in doc.blocks if b.kind is StructuredBlockKind.PARAGRAPH]
    assert len(paragraphs) == 2
    assert paragraphs[0].text.startswith("Hola mundo")
    assert all(b.char_range is not None and b.char_range.length > 0 for b in doc.blocks)
    assert doc.section_count == 0
    assert doc.content_hash


def test_parser_registry_resolves_and_normalizes_extensions() -> None:
    assert isinstance(get_parser("pdf"), PdfParser)
    assert isinstance(get_parser("md"), TextParser)
    assert isinstance(get_parser(".PDF"), PdfParser)
    assert get_parser("txt") is get_parser("md")  # mismo parser compartido
    assert get_parser("exe") is None
    assert "pdf" in supported_extensions()
    assert "md" in supported_extensions()


def test_pdf_parser_extracts_pages_blocks_bbox_and_reading_order() -> None:
    data = _tiny_pdf(["Section 5.2 Commission", "Commission rate is five percent."])
    doc = PdfParser().parse(
        data,
        organization_id=uuid4(),
        external_id="manual.pdf",
        source_name="manual.pdf",
        mime_type="application/pdf",
    )
    doc.check_consistency()

    assert doc.mime_type == "application/pdf"
    assert doc.page_count == 1
    assert doc.pages[0].page_number == 1
    assert "Commission" in doc.pages[0].text
    assert doc.block_count == len(doc.blocks) == 2
    assert all(b.page == 1 for b in doc.blocks)
    assert all(b.bbox is not None and b.bbox.page == 1 for b in doc.blocks)

    # orden de lectura: la línea superior (y0 menor) va primero
    assert "Section 5.2 Commission" in doc.blocks[0].text
    assert "five percent" in doc.blocks[1].text


def test_pdf_parser_rejects_invalid_pdf() -> None:
    with pytest.raises(StructuredParserError):
        PdfParser().parse(
            b"not a pdf",
            organization_id=uuid4(),
            external_id="broken.pdf",
        )


def _build_docx_bytes() -> bytes:
    import io

    from docx import Document as DocxDocument

    buffer = io.BytesIO()
    doc = DocxDocument()
    doc.add_heading("Manual Operaciones", level=0)
    doc.add_heading("Compensación", level=1)
    doc.add_paragraph("Comisión 5% para contratos vigentes.")
    doc.add_heading("Comisiones", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Clave"
    table.cell(0, 1).text = "Valor"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "5%"
    doc.add_paragraph("Cláusula con viñeta", style="List Bullet")
    doc.save(buffer)
    return buffer.getvalue()


def test_docx_parser_builds_blocks_sections_and_tables() -> None:
    from src.knowledge.structure import DocxParser

    doc = DocxParser().parse(
        _build_docx_bytes(),
        organization_id=uuid4(),
        external_id="manual.docx",
    )
    doc.check_consistency()

    kinds = {b.kind for b in doc.blocks}
    assert StructuredBlockKind.TITLE in kinds
    assert StructuredBlockKind.HEADING in kinds
    assert StructuredBlockKind.PARAGRAPH in kinds
    assert StructuredBlockKind.TABLE in kinds

    assert doc.section_count == 2
    assert doc.sections[0].heading == "Compensación"
    assert doc.sections[0].parent_id is None
    assert doc.sections[1].heading == "Comisiones"
    assert doc.sections[1].parent_id == doc.sections[0].id

    assert doc.table_count == 1
    assert doc.tables[0].headers == ("Clave", "Valor")
    assert doc.tables[0].rows == (("A", "5%"),)
    assert doc.title == "Manual Operaciones"
    assert doc.pages == ()


_HTML = """<!doctype html><html lang="es"><head><title>Política</title></head>
<body>
<h1>Operaciones</h1>
<p>Comisión 5% para contratos vigentes.</p>
<h2>5.2 Comisiones</h2>
<table><tr><th>Clave</th><th>Valor</th></tr><tr><td>A</td><td>5%</td></tr></table>
<ul><li>cláusula uno</li><li>cláusula dos</li></ul>
<img src="logo.png" alt="Logo Zent">
</body></html>
"""


def test_html_parser_builds_blocks_sections_tables_and_figures() -> None:
    from src.knowledge.structure import HtmlParser

    doc = HtmlParser().parse(
        _HTML.encode("utf-8"),
        organization_id=uuid4(),
        external_id="politica.html",
    )
    doc.check_consistency()

    assert doc.title == "Política"
    assert doc.language == "es"
    assert doc.section_count == 2
    assert [s.heading for s in doc.sections] == ["Operaciones", "5.2 Comisiones"]
    assert doc.sections[1].parent_id == doc.sections[0].id

    assert doc.table_count == 1
    assert doc.tables[0].headers == ("Clave", "Valor")
    assert doc.figure_count == 1
    assert doc.figures[0].caption == "Logo Zent"
    kinds = {b.kind for b in doc.blocks}
    assert StructuredBlockKind.LIST in kinds
    assert StructuredBlockKind.HEADING in kinds
    assert doc.pages == ()


def test_registry_covers_pdf_text_docx_html() -> None:
    from src.knowledge.structure import DocxParser, HtmlParser

    assert isinstance(get_parser("pdf"), PdfParser)
    assert isinstance(get_parser("md"), TextParser)
    assert isinstance(get_parser("docx"), DocxParser)
    assert isinstance(get_parser("html"), HtmlParser)
    assert get_parser("htm") is get_parser("html")
    assert get_parser("exe") is None
