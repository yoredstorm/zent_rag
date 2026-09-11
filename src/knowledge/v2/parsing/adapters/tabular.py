# =============================================================================
# CSV / Excel / JSON adapters
# =============================================================================
from __future__ import annotations

import csv
import io
import json

from src.core.domain.knowledge_v2 import (
    PAGE_ABSENT_FORMAT_HAS_NO_PAGES,
    SECTION_ABSENT_NO_HEADINGS,
    StructuredBlockKind,
    StructuredDocument,
)
from src.knowledge.v2.parsing.adapters.base import ParseError, StructuredAdapter
from src.knowledge.v2.parsing.adapters.text import decode_text
from src.knowledge.v2.parsing.blocks import assign_char_spans, make_block
from src.knowledge.v2.parsing.context import ParseContext, assemble_document


def _section_meta(path: tuple[str, ...]) -> dict:
    if path:
        return {}
    return {"section_absent_reason": SECTION_ABSENT_NO_HEADINGS}


def markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [list(r) + [""] * (width - len(r)) for r in rows]
    header = [c.strip() or f"col_{i}" for i, c in enumerate(norm[0])]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in norm[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


class CsvAdapter(StructuredAdapter):
    name = "csv"
    mime_type = "text/csv"

    def parse(self, ctx: ParseContext) -> StructuredDocument:
        text = decode_text(ctx.data)
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        rows = [[cell.strip() for cell in row] for row in reader if any(c.strip() for c in row)]
        if not rows:
            raise ParseError(f"CSV produced no rows: {ctx.filename}")
        path = (ctx.stem,)
        table = markdown_table(rows)
        blocks = [
            make_block(
                kind=StructuredBlockKind.TABLE,
                text=table,
                order=0,
                page=None,
                page_absent_reason=PAGE_ABSENT_FORMAT_HAS_NO_PAGES,
                heading_path=path,
                metadata={"format": "csv", "row_count": len(rows)},
            )
        ]
        blocks = assign_char_spans(blocks)
        return assemble_document(ctx, blocks, mime_type=self.mime_type, adapter=self.name)


class ExcelAdapter(StructuredAdapter):
    name = "xlsx"
    mime_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def parse(self, ctx: ParseContext) -> StructuredDocument:
        try:
            import openpyxl
        except ImportError as exc:
            raise ParseError("openpyxl is required for Excel V2 parsing") from exc
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(ctx.data), read_only=True, data_only=True)
        except Exception as exc:
            raise ParseError(f"Excel parse failed for {ctx.filename}: {exc}") from exc
        blocks: list = []
        order = 0
        try:
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                rows: list[list[str]] = []
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if cell is None else str(cell) for cell in row]
                    if any(v.strip() for v in values):
                        rows.append(values)
                if not rows:
                    continue
                heading = make_block(
                    kind=StructuredBlockKind.HEADING,
                    text=sheet_name,
                    order=order,
                    page=None,
                    page_absent_reason=PAGE_ABSENT_FORMAT_HAS_NO_PAGES,
                    heading_path=(sheet_name,),
                    metadata={"sheet": sheet_name},
                )
                order += 1
                table = make_block(
                    kind=StructuredBlockKind.TABLE,
                    text=markdown_table(rows),
                    order=order,
                    page=None,
                    page_absent_reason=PAGE_ABSENT_FORMAT_HAS_NO_PAGES,
                    heading_path=(sheet_name,),
                    metadata={"sheet": sheet_name, "row_count": len(rows)},
                )
                order += 1
                blocks.extend([heading, table])
        finally:
            workbook.close()
        if not blocks:
            raise ParseError(f"Excel produced no cells: {ctx.filename}")
        blocks = assign_char_spans(blocks)
        return assemble_document(ctx, blocks, mime_type=self.mime_type, adapter=self.name)


class JsonAdapter(StructuredAdapter):
    name = "json"
    mime_type = "application/json"

    def parse(self, ctx: ParseContext) -> StructuredDocument:
        text = decode_text(ctx.data)
        try:
            payload = json.loads(text)
        except json.JSONDecodeError:
            path: tuple[str, ...] = ()
            blocks = [
                make_block(
                    kind=StructuredBlockKind.CODE,
                    text=text,
                    order=0,
                    page=None,
                    page_absent_reason=PAGE_ABSENT_FORMAT_HAS_NO_PAGES,
                    heading_path=path,
                    metadata=_section_meta(path) | {"json_parse": "invalid"},
                )
            ]
            return assemble_document(
                ctx, assign_char_spans(blocks), mime_type=self.mime_type, adapter=self.name
            )

        blocks = []
        order = 0

        def emit(kind: StructuredBlockKind, body: str, path: tuple[str, ...]) -> None:
            nonlocal order
            blocks.append(
                make_block(
                    kind=kind,
                    text=body,
                    order=order,
                    page=None,
                    page_absent_reason=PAGE_ABSENT_FORMAT_HAS_NO_PAGES,
                    heading_path=path,
                    metadata=_section_meta(path),
                )
            )
            order += 1

        if isinstance(payload, dict):
            for key, value in payload.items():
                path = (str(key),)
                emit(StructuredBlockKind.HEADING, str(key), path)
                if isinstance(value, list) and value and isinstance(value[0], dict):
                    keys = list(dict.fromkeys(k for row in value if isinstance(row, dict) for k in row))
                    rows = [keys] + [[str(row.get(k, "")) for k in keys] for row in value if isinstance(row, dict)]
                    emit(StructuredBlockKind.TABLE, markdown_table(rows), path)
                elif isinstance(value, (dict, list)):
                    emit(StructuredBlockKind.CODE, json.dumps(value, ensure_ascii=False, indent=2), path)
                else:
                    emit(StructuredBlockKind.PARAGRAPH, f"{key}: {value}", path)
        elif isinstance(payload, list) and payload and isinstance(payload[0], dict):
            keys = list(dict.fromkeys(k for row in payload if isinstance(row, dict) for k in row))
            rows = [keys] + [[str(row.get(k, "")) for k in keys] for row in payload if isinstance(row, dict)]
            emit(StructuredBlockKind.TABLE, markdown_table(rows), (ctx.stem,))
        else:
            emit(
                StructuredBlockKind.CODE,
                json.dumps(payload, ensure_ascii=False, indent=2),
                (ctx.stem,),
            )
        if not blocks:
            emit(StructuredBlockKind.CODE, text, ())
        return assemble_document(
            ctx, assign_char_spans(blocks), mime_type=self.mime_type, adapter=self.name
        )
