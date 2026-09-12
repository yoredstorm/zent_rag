# =============================================================================
# Knowledge V2 — DOCX parser (python-docx)
# =============================================================================
# Recorre el body de Word en orden real (párrafos + tablas intercaladas),
# detecta headings por estilo ('Heading N'), listas y tablas. Sin páginas ni
# bbox (formato de flujo); el árbol de secciones se construye por nivel.
# =============================================================================
from __future__ import annotations

import io
import re
from uuid import UUID, uuid4, uuid5

from src.core.domain.catalog import CatalogProvenance
from src.core.domain.knowledge_v2 import (
    DocumentSection,
    DocumentTable,
    KnowledgeObjectStatus,
    StructuredBlock,
    StructuredBlockKind,
    StructuredDocument,
)
from src.knowledge.structure.base import (
    StructuredParser,
    StructuredParserError,
    content_hash,
    token_count,
)

_NS = UUID("d1e4f6a8-9b2c-4d5e-8f7a-6c3b9e1a2f40")
_HEADING_STYLE_RE = re.compile(r"heading\s*(\d)", re.IGNORECASE)
_LIST_STYLE_RE = re.compile(r"(List\s*(?:Bullet|Number))", re.IGNORECASE)


def _stable_document_id(organization_id: UUID, source_id: UUID | None, external_id: str) -> UUID:
    return uuid5(_NS, f"v2:{organization_id}:{source_id}:{external_id}")


class DocxParser(StructuredParser):
    """docx → StructuredDocument (bloques + secciones + tablas en orden)."""

    kind = "docx"
    mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def parse(
        self,
        data: bytes,
        *,
        organization_id: UUID,
        external_id: str,
        source_id: UUID | None = None,
        workspace_id: UUID | None = None,
        source_name: str = "document",
        mime_type: str | None = None,
    ) -> StructuredDocument:
        try:
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as exc:  # pragma: no cover
            raise StructuredParserError("python-docx is not installed") from exc

        try:
            document_docx = Document(io.BytesIO(data))
        except Exception as exc:
            raise StructuredParserError(f"DOCX could not be opened ({source_name}): {exc}") from exc

        blocks: list[StructuredBlock] = []
        tables: list[DocumentTable] = []
        sections: list[DocumentSection] = []
        section_stack: list[tuple[int, DocumentSection]] = []
        order = 0
        doc_id = _stable_document_id(organization_id, source_id, external_id)

        def next_order() -> int:
            nonlocal order
            value = order
            order += 1
            return value

        def emit_block(kind: StructuredBlockKind, text_value: str, heading_level: int = 0) -> None:
            bs = StructuredBlock(
                kind=kind,
                text=text_value,
                order=next_order(),
                token_count=token_count(text_value),
                content_hash=content_hash(text_value),
            )
            blocks.append(bs)
            if kind is StructuredBlockKind.HEADING:
                _attach_heading_section(
                    bs, heading_level, organization_id, workspace_id, source_id,
                    doc_id, sections, section_stack,
                )

        body_iter = iter(document_docx.element.body)
        table_iter = iter(document_docx.tables)
        for child in body_iter:
            tag = child.tag.split("}")[-1]
            if tag == "p":
                paragraph = Paragraph(child, document_docx)
                text = paragraph.text.strip()
                if not text:
                    continue
                style_name = (paragraph.style.name if paragraph.style else "") or ""
                heading_match = _HEADING_STYLE_RE.search(style_name)
                if heading_match:
                    emit_block(StructuredBlockKind.HEADING, text, int(heading_match.group(1)))
                elif _LIST_STYLE_RE.search(style_name) or text.startswith(("•", "- ", "1. ")):
                    emit_block(StructuredBlockKind.LIST, text)
                elif style_name.lower() == "title":
                    emit_block(StructuredBlockKind.TITLE, text)
                else:
                    emit_block(StructuredBlockKind.PARAGRAPH, text)
            elif tag == "tbl":
                try:
                    table = next(table_iter)
                except StopIteration:  # pragma: no cover - estructura inconsistente
                    continue
                if isinstance(table, Table):
                    parsed_table = _build_table(table, doc_id, organization_id, workspace_id, source_id, next_order())
                    tables.append(parsed_table)
                    emit_block(StructuredBlockKind.TABLE, " | ".join(parsed_table.headers))

        document_text = "\n".join(b.text for b in blocks)
        return StructuredDocument(
            id=doc_id,
            organization_id=organization_id,
            workspace_id=workspace_id,
            source_id=source_id,
            external_id=external_id,
            title=_title_from_docx(document_docx) or source_name,
            content_hash=content_hash(document_text),
            mime_type=mime_type or self.mime_type,
            language=None,
            blocks=tuple(blocks),
            pages=(),
            sections=tuple(sections),
            tables=tuple(tables),
            figures=(),
            provenance=CatalogProvenance.OBSERVED,
            status=KnowledgeObjectStatus.OBSERVED,
        )


def _title_from_docx(document_docx) -> str | None:
    core = getattr(document_docx, "core_properties", None)
    if core is not None and getattr(core, "title", ""):
        return core.title
    for paragraph in document_docx.paragraphs:
        if paragraph.text.strip() and paragraph.style and "title" in (paragraph.style.name or "").lower():
            return paragraph.text.strip()
    return None


def _attach_heading_section(
    block: StructuredBlock,
    heading_level: int,
    organization_id: UUID,
    workspace_id: UUID | None,
    source_id: UUID | None,
    document_id: UUID,
    sections: list[DocumentSection],
    section_stack: list[tuple[int, DocumentSection]],
) -> None:
    while section_stack and section_stack[-1][0] >= heading_level:
        section_stack.pop()
    ancestors = [s[1] for s in section_stack]
    path = tuple(a.heading for a in ancestors) + (block.text,)
    parent_id = ancestors[-1].id if ancestors else None
    section = DocumentSection(
        id=uuid4(),
        document_id=document_id,
        organization_id=organization_id,
        workspace_id=workspace_id,
        source_id=source_id,
        section_path=path,
        heading=block.text,
        depth=heading_level - 1,
        parent_id=parent_id,
        order=block.order,
        block_ids=(block.id,),
        text=block.text,
        token_count=token_count(block.text),
        content_hash=content_hash(block.text),
    )
    sections.append(section)
    section_stack.append((heading_level, section))


def _build_table(
    table,
    document_id: UUID,
    organization_id: UUID,
    workspace_id: UUID | None,
    source_id: UUID | None,
    order: int,
) -> DocumentTable:
    rows: list[tuple[str, ...]] = []
    for row in table.rows:
        cells = tuple((cell.text or "").strip() for cell in row.cells)
        if any(cells):
            rows.append(cells)
    headers = rows[0] if rows else ()
    body = tuple(rows[1:])
    rendered = "\n".join(" | ".join(r) for r in rows)
    return DocumentTable(
        id=uuid4(),
        document_id=document_id,
        organization_id=organization_id,
        workspace_id=workspace_id,
        source_id=source_id,
        headers=headers,
        rows=body,
        token_count=token_count(rendered),
        content_hash=content_hash(rendered),
        metadata={"docx_order": order},
    )
