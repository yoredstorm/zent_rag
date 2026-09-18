# =============================================================================
# Data Onboarding — perfil de archivos (CSV/Excel/documentos)
# =============================================================================
from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from uuid import UUID

import src.knowledge.normalize  # noqa: F401 — registra PDF/docx/html
from src.knowledge.normalize.base import get_normalizer
from src.knowledge.storage import resolve_path


def profile_upload(
    organization_id: UUID,
    object_key: str,
    source_type: str,
    filename: str,
) -> dict:
    path = resolve_path(organization_id, object_key)
    if not path.exists():
        return {"kind": source_type, "filename": filename, "error": "file_missing"}
    data = path.read_bytes()
    if source_type in ("csv",):
        return _profile_csv(data, filename)
    if source_type == "excel":
        return _profile_excel(path, filename)
    return _profile_document(data, filename)


def _profile_csv(data: bytes, filename: str) -> dict:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return {"kind": "spreadsheet", "filename": filename, "row_count": 0, "columns": []}
    header = [str(h).strip() or f"col_{i}" for i, h in enumerate(rows[0])]
    body = rows[1:]
    columns = []
    for i, name in enumerate(header):
        values = [r[i] if i < len(r) else "" for r in body]
        nonempty = [v for v in values if v not in (None, "")]
        null_ratio = round(1 - (len(nonempty) / max(len(values), 1)), 3)
        distinct = len(set(nonempty))
        inferred = _infer(nonempty)
        columns.append(
            {
                "physical_name": name,
                "inferred_type": inferred,
                "null_ratio": null_ratio,
                "distinct_values": distinct,
                "possible_meanings": _guess_meanings(name),
                "uncertain": len(_guess_meanings(name)) > 1 and _guess_meanings(name)[0][1] < 70,
            }
        )
    entity = _guess_entity(filename, header)
    return {
        "kind": "spreadsheet",
        "filename": filename,
        "row_count": len(body),
        "likely_entity": entity,
        "columns": columns,
        "interpretation": [
            {
                "physical": c["physical_name"],
                "business": c["possible_meanings"][0][0] if c["possible_meanings"] else c["physical_name"],
            }
            for c in columns
        ],
    }


def _profile_excel(path: Path, filename: str) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"kind": "spreadsheet", "filename": filename, "sheets": [], "error": "openpyxl_missing"}
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    primary_columns: list[dict] = []
    for sheet_index, sheet_name in enumerate(wb.sheetnames):
        ws = wb[sheet_name]
        rows = [list(row) for _, row in zip(range(6), ws.iter_rows(values_only=True), strict=False)]
        header = [str(c) if c is not None else f"col_{i}" for i, c in enumerate(rows[0])] if rows else []
        sheets.append({"name": sheet_name, "columns": header})
        if sheet_index == 0 and header:
            # La PRIMERA hoja alimenta `columns` del understanding (lo usa el
            # readiness y el catálogo). Muestras de las filas siguientes para
            # inferir tipo y null_ratio sin leer todo el archivo.
            sample_rows = rows[1:]
            for idx, name in enumerate(header):
                values = [
                    row[idx]
                    for row in sample_rows
                    if idx < len(row) and row[idx] is not None and str(row[idx]).strip() != ""
                ]
                rendered = [str(value) for value in values]
                inferred = _infer(rendered) if rendered else "text"
                null_ratio = round(1 - (len(rendered) / max(len(sample_rows), 1)), 3)
                primary_columns.append(
                    {
                        "physical_name": name,
                        "inferred_type": inferred,
                        "null_ratio": null_ratio,
                        "distinct_values": len(set(rendered)),
                        "possible_meanings": _guess_meanings(name),
                        "uncertain": False,
                    }
                )
    wb.close()
    result = {"kind": "spreadsheet", "filename": filename, "sheets": sheets}
    if primary_columns:
        result["columns"] = primary_columns
        result["likely_entity"] = _guess_entity(filename, [c["physical_name"] for c in primary_columns])
        result["interpretation"] = [
            {
                "physical": column["physical_name"],
                "business": (
                    column["possible_meanings"][0][0]
                    if column["possible_meanings"]
                    else column["physical_name"]
                ),
            }
            for column in primary_columns
        ]
    return result


def _profile_document(data: bytes, filename: str) -> dict:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "txt"
    text = ""
    pages = None
    try:
        normalizer = get_normalizer(ext)
        if normalizer is not None:
            text = normalizer.normalize(data, source_name=filename)
        else:
            text = data[:8000].decode("utf-8", errors="replace")
    except Exception:
        text = data[:8000].decode("utf-8", errors="replace")
    headings = re.findall(r"^#{1,3}\s+(.+)$", text, flags=re.M)
    if not headings:
        headings = [
            ln.strip()
            for ln in text.splitlines()
            if ln.strip() and ln.strip() == ln.strip().upper() and 3 < len(ln.strip()) < 80
        ][:12]
    dates = re.findall(r"\b20\d{2}-\d{2}-\d{2}\b", text)
    topics = headings[:8]
    title = headings[0] if headings else filename
    return {
        "kind": "document",
        "filename": filename,
        "title": title[:160],
        "pages": pages,
        "headings": headings[:20],
        "topics": topics,
        "dates": dates[:10],
        "document_type": _guess_doc_type(text, filename),
        "excerpt": text[:600],
    }


def _infer(values: list[str]) -> str:
    if not values:
        return "text"
    if all(_is_int(v) for v in values):
        return "integer"
    if all(_is_float(v) for v in values):
        return "number"
    return "text"


def _is_int(value: str) -> bool:
    try:
        int(value)
        return True
    except ValueError:
        return False


def _is_float(value: str) -> bool:
    try:
        float(value.replace(",", "."))
        return True
    except ValueError:
        return False


def _guess_entity(filename: str, header: list[str]) -> str:
    joined = " ".join([filename] + header).lower()
    for name, keys in (
        ("Products", ("product", "sku", "precio", "stock", "codigo")),
        ("Customers", ("customer", "cliente", "email")),
        ("Sales", ("sale", "venta", "amount", "monto")),
    ):
        if any(k in joined for k in keys):
            return name
    return "Records"


def _guess_meanings(name: str) -> list[tuple[str, int]]:
    n = name.lower().strip()
    mapping = [
        (("codigo", "sku", "code", "id"), "Product Code", 90),
        (("descrip", "desc", "nombre", "name", "title"), "Product Description", 88),
        (("precio", "price", "fare", "pvp"), "Unit Price", 86),
        (("cost", "costo"), "Unit Cost", 80),
        (("stock", "qty", "cantidad"), "Current Stock", 84),
        (("fecha", "date"), "Date", 82),
        (("val",), "Amount", 42),
    ]
    hits: list[tuple[str, int]] = []
    for keys, label, score in mapping:
        if any(k in n for k in keys):
            hits.append((label, score))
    if n == "val":
        return [("Unit Price", 42), ("Unit Cost", 31), ("Amount", 17)]
    return hits or [(name.replace("_", " ").title(), 50)]


def _guess_doc_type(text: str, filename: str) -> str:
    blob = (filename + " " + text[:2000]).lower()
    if any(k in blob for k in ("vacation", "vacacion", "hr policy", "remote work")):
        return "HR Policy"
    if "contract" in blob or "contrato" in blob:
        return "Contract"
    if "manual" in blob:
        return "Manual"
    return "Document"
